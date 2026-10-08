#!/usr/bin/env python3
"""Build the shared Tariff-Impact Model template (Research + Standard Cost).

Run from the repo root (build/build.py calls it):

    python3 build/make_tariff_model.py            # writes dist/Tariff-Impact-Model-Template.xlsx
    python3 build/make_tariff_model.py --out X    # writes somewhere else

What it makes: a LIVE Excel workbook. Every result is an Excel formula, not a typed value.
Tabs: README, Action (Research fills: the WHAT), Items (Standard Cost fills: the $ inputs),
Calc (formulas only), Summary (formulas only), Review (maker-checker sign-off log).

What the model handles (template v2.0):
- Several rate lines per action, keyed by rate group (from the HTS scope list), origin and USMCA status.
- A partial duty base (PART): the rate applies to a share of the customs value (metal content,
  non-US content), entered per item on Items.
- Partial USMCA relief (a rate line for USMCA-qualifying goods), or full exemption.
- Non-stacking rules: an other additional duty that the action switches off, per rate line or per item.
- Research determinations keyed by item AND import country, trimmed, with a matched-count check.
- Duplicate item and import-country rows, unit-of-measure and USMCA-origin checks.
- Margin % on items with a sell price only; impact on unpriced items is still counted and shown.
- A quantity period with a label (Annual by default).
- Every tab is protected without a password; input (yellow) cells are unlocked.

Rules this script follows:
- Formulas work in Excel 2010+ and LibreOffice Calc: no dynamic arrays, LET, LAMBDA, XLOOKUP,
  IFS or wildcard criteria. Lookups use INDEX/MATCH, COUNTIF(S), SUMIFS and SUMPRODUCT
  (SUMPRODUCT(MAX(...)) evaluates its array without array entry in both).
- The sample rows are SYNTHETIC (a fictional action and invented items) and say so on the sheet.
- The output is byte-for-byte deterministic: fixed document timestamps and fixed zip entry times.

Explained tab by tab in docs/templates/domain/tariff-impact-model-guide.md.
"""

import argparse
import datetime as dt
import re
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "dist" / "Tariff-Impact-Model-Template.xlsx"

TEMPLATE_VERSION = "2.0"
TEMPLATE_DATE = dt.datetime(2026, 10, 8, 0, 0, 0)
FIRST = 5             # first data row on Action lists, Items and Calc
ITEM_ROWS = 250       # capacity of Items / Calc
HTS_ROWS = 300        # capacity of the HTS scope list and the HTS exclusion list
ORIGIN_ROWS = 50      # capacity of the origin list (and of the Summary origin table)
DET_ROWS = 100        # capacity of the item-level determinations list
RATE_ROWS = 30        # capacity of the rate lines table
CHAPTER_ROWS = 10     # yellow chapter rows on Summary
LAST = FIRST + ITEM_ROWS - 1
HTS_LAST = FIRST + HTS_ROWS - 1
ORIGIN_LAST = FIRST + ORIGIN_ROWS - 1
DET_LAST = FIRST + DET_ROWS - 1
RATE_LAST = FIRST + RATE_ROWS - 1
assert HTS_ROWS < 1000 and RATE_ROWS < 100  # the scope and rate scores pack the row index into the low digits

NAVY = "1F3864"
FONT = "Calibri"
INPUT_YELLOW = "FFF2CC"
CALC_GREY = "F2F2F2"
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
THIN = Side(style="thin", color="BFBFBF")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ------------------------------------------------------------------ styles ----

STYLE_SPECS = {
    # name: (font kwargs, fill colour or None, number format, wrap, border)
    "TIM Title": (dict(size=16, bold=True, color=NAVY), None, "General", False, False),
    "TIM Subtitle": (dict(italic=True, color="595959"), None, "General", False, False),
    "TIM Warning": (dict(bold=True, color="C00000"), None, "General", False, False),
    "TIM Section": (dict(size=12, bold=True, color=NAVY), None, "General", False, False),
    "TIM Header": (dict(bold=True, color="FFFFFF"), NAVY, "General", True, True),
    "TIM Label": (dict(bold=True, size=10), None, "General", True, True),
    "TIM Text": (dict(size=10), None, "General", True, False),
    "TIM Note": (dict(size=9, italic=True, color="595959"), None, "General", True, True),
    "TIM Input": (dict(size=10), INPUT_YELLOW, "General", True, True),
    "TIM Input Text": (dict(size=10), INPUT_YELLOW, "@", True, True),
    "TIM Input Pct": (dict(size=10), INPUT_YELLOW, "0.0%", False, True),
    "TIM Input Money": (dict(size=10), INPUT_YELLOW, "#,##0.00", False, True),
    "TIM Input Qty": (dict(size=10), INPUT_YELLOW, "#,##0", False, True),
    "TIM Input Date": (dict(size=10), INPUT_YELLOW, "yyyy-mm-dd", False, True),
    "TIM Calc": (dict(size=10), CALC_GREY, "General", False, True),
    "TIM Calc Pct": (dict(size=10), CALC_GREY, "0.0%", False, True),
    "TIM Calc Money": (dict(size=10), CALC_GREY, "#,##0.00", False, True),
    "TIM Calc Unit": (dict(size=10), CALC_GREY, "#,##0.0000", False, True),
    "TIM Calc Whole": (dict(size=10), CALC_GREY, "#,##0", False, True),
    "TIM Calc Date": (dict(size=10), CALC_GREY, "yyyy-mm-dd", False, True),
    "TIM Calc Key": (dict(size=8, color="808080"), CALC_GREY, "0.000000000", False, True),
}


def add_styles(wb):
    for name, (font, fill, fmt, wrap, border) in STYLE_SPECS.items():
        st = NamedStyle(name=name)
        st.font = Font(name=FONT, **font)
        if fill:
            st.fill = PatternFill("solid", fgColor=fill)
        st.number_format = fmt
        st.alignment = Alignment(wrap_text=wrap, vertical="top")
        if border:
            st.border = GRID
        # Input cells stay editable when the sheet is protected; everything else is locked.
        st.protection = Protection(locked=not name.startswith("TIM Input"))
        wb.add_named_style(st)


def put(ws, ref, value, style):
    cell = ws[ref]
    cell.value = value
    cell.style = style
    return cell


def title(ws, text, subtitle, warning=None):
    put(ws, "A1", text, "TIM Title")
    put(ws, "A2", subtitle, "TIM Subtitle")
    if warning:
        put(ws, "A3", warning, "TIM Warning")


def headers(ws, row, col, names):
    for i, label in enumerate(names):
        put(ws, f"{get_column_letter(col + i)}{row}", label, "TIM Header")


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def list_dv(ws, values, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "Pick a value from the list.", "Not allowed", True
    if prompt:
        dv.prompt, dv.showInputMessage = prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def number_dv(ws, kind, rng, low, high, message):
    dv = DataValidation(type=kind, operator="between", formula1=str(low), formula2=str(high), allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = message, "Check the value", True
    ws.add_data_validation(dv)
    dv.add(rng)


def length_dv(ws, rng, length, message):
    dv = DataValidation(type="textLength", operator="equal", formula1=str(length), allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = message, "Check the value", True
    ws.add_data_validation(dv)
    dv.add(rng)


def name(wb, label, ref):
    wb.defined_names[label] = DefinedName(label, attr_text=ref)


def protect(ws, *, sort=False, filter_=False):
    """Protect the sheet without a password. Yellow input cells are unlocked by their style."""
    p = ws.protection
    p.sheet = True
    p.formatColumns = False      # False = allowed: column widths can change
    p.formatRows = False
    p.sort = not sort            # in OOXML, True means "blocked"
    p.autoFilter = not filter_
    # Row and column insert/delete stay blocked: they would break the row-by-row link Items -> Calc.


# ------------------------------------------------------------------ README ----

README_ROWS = [
    ("section", "What this workbook is"),
    ("", "The one shared model for a tariff action. Research (Hire A) records the WHAT on the Action tab. "
         "Standard Cost (Hire B) loads the items on the Items tab. Calc and Summary work out the $ impact with live formulas."),
    ("", "Boundary rule: Research owns scope, HTS, origins, rates, rate lines, duty base rules, dates, stacking and item-level "
         "determinations (from the checked Tariff-Impact Brief). Standard Cost owns the item data, landed cost, margin and the "
         "Oracle cost update, and never decides scope. Nobody edits the other track's tab; ask on the Jira story instead."),
    ("", "Guide (every tab and formula explained, with sample results and drills): Tariff-Impact Model Guide. Brief: Tariff-Impact "
         "Brief Template. Both are in the team OneNote and the onboarding folder."),
    ("section", "Colour key and protection"),
    ("Yellow", "Input cell. Type or paste values here only (Paste Special > Values)."),
    ("Grey", "Formula cell. Locked. If a formula looks wrong, raise it with the Manager; never patch it with a typed number."),
    ("Green / red / amber", "Status from conditional formatting: green = in scope or check passed; red = a CHECK that must be cleared "
                            "before review; amber = a NOTE that needs one line in the covering note."),
    ("Protection", "Every tab is protected without a password, so only yellow cells can be changed and rows cannot be inserted or "
                   "deleted (that would break the row-by-row link between Items and Calc). Unprotect only with the Manager's "
                   "agreement, and protect again before review."),
    ("section", "How to use it (one copy per tariff action and importing country)"),
    ("1", "Copy this template. Name the copy TIM_[Action ID]_[import country]_v0.1.xlsx and save it in [COMPANY: tariff model folder]. "
          "Link it on the Jira story; never email copies."),
    ("2", "Clear the SAMPLE rows on Action (lists, determinations, rate lines), Items and Review before real use: select the cells "
          "and press Delete. Never delete rows."),
    ("3", "Research (maker): fill the Action tab from the checked brief: fields in column B, the HTS scope list with rate groups, "
          "the exclusion list, the origin list, item-level determinations and rate lines."),
    ("4", "Research checker (Senior or Manager): compare the Action tab with the brief line by line. Log the review on the Review tab."),
    ("5", "Standard Cost (maker): set the quantity period and currency (Items K1:K2), then paste the item extract into Items as values "
          "([COMPANY: Oracle cost extract or warehouse query], with the extract date in column N). One row per item and import country, "
          "all values per the Oracle primary unit of measure in column G. Keep HTS, item numbers and origin codes as text."),
    ("6", "Clear every CHECK in Calc column AJ and on the Action status columns. Summary B30 must read YES before review."),
    ("7", "Standard Cost checker (Manager): recalculate 3 items by hand, confirm the item count and totals against the extract, "
          "confirm the Summary reconciles. Log the review on the Review tab."),
    ("8", "Release: only a version with a checker sign-off on the Review tab goes to the business. Any change after sign-off is a new "
          "version and a new review."),
    ("section", "Maker-checker rule"),
    ("", "Every model released to the business gets a documented second-person review (Senior or Manager): name, date, version, findings. "
         "This applies after day 90 too. The maker never signs their own review."),
    ("section", "Limits (read before you trust a number)"),
    ("", "Ad valorem rates only. Specific or compound rates (per kg, per unit) need an ad valorem equivalent agreed with Research and the checker."),
    ("", "Landed cost change = duty change only. Fees (for example US MPF), freight and brokerage are assumed unchanged. Duty-based taxes "
         "(Canada GST, Mexico IVA) are left out because they are normally recoverable; confirm with Finance."),
    ("", "The impact applies the action to the whole quantity period on Items (Annual by default). It does not apply the effective date, "
         "in-transit relief, stock already on hand, the end date or supplier price changes."),
    ("", "Duty paid is assumed. Duty deferral or recovery (Mexico IMMEX temporary imports, US drawback) is not modelled: exclude or "
         "adjust those flows and say so in the covering note."),
    ("", "Current duty rate and current standard cost come from the Items tab as given. The model assumes the current standard cost "
         "already includes duty at the current rate."),
    ("", f"Capacity: {ITEM_ROWS} item rows; {HTS_ROWS} HTS prefixes and {HTS_ROWS} exclusions; {ORIGIN_ROWS} origins; "
         f"{DET_ROWS} determinations; {RATE_ROWS} rate lines. To extend, see the guide."),
    ("section", "Template version"),
    ("", f"Template v{TEMPLATE_VERSION}, built {TEMPLATE_DATE:%Y-%m-%d} by build/make_tariff_model.py. Owner: BSA Manager. "
         "Change the generator script, not a copy of the workbook."),
]


def build_readme(wb):
    ws = wb.active
    ws.title = "README"
    title(ws, "Tariff-Impact Model (template)",
          "Shared by Research (the WHAT) and Standard Cost (the $). One copy per tariff action and importing country.",
          "All sample rows in this template are SYNTHETIC and use a FICTIONAL action. Clear them before real use.")
    r = 5
    for label, text in README_ROWS:
        if label == "section":
            r += 1
            put(ws, f"A{r}", text, "TIM Section")
        else:
            put(ws, f"A{r}", label, "TIM Label")
            put(ws, f"B{r}", text, "TIM Text")
            ws.row_dimensions[r].height = 15 * max(1, -(-len(text) // 120))
        r += 1
    widths(ws, {"A": 16, "B": 130})
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = "808080"
    protect(ws)


# ------------------------------------------------------------------ Action ----

# (label, defined name, sample value, style, note). Rows 5 to 24, in this order.
# Rows B8 and B10 to B19 are referenced by docs/primers/current-tariff-landscape.md (section 2.7): keep them in place.
ACTION_FIELDS = [
    ("Action ID", "act_ID", "TA-US-FIC-01", "TIM Input Text",
     "Brief Part 1. Same ID in the brief, this file name and the Jira story."),
    ("Action name", "act_Name",
     "FICTIONAL EXAMPLE: US additional duty on iron or steel fasteners (heading 7318) from Country X", "TIM Input",
     "Brief Part 1."),
    ("Importing country (US, CA or MX)", "act_Country", "US", "TIM Input Text",
     "Brief Part 1. One model per importing country."),
    ("Legal authority category", "act_Authority", "Section 301", "TIM Input",
     "Brief Part 3. Category only; the citation goes in B10."),
    ("Tariff-Impact Brief ID and version", "act_Brief", "TIB-TA-US-FIC-01 v1.0", "TIM Input Text",
     "Brief Part 1. The model must match a checked brief version."),
    ("Primary legal source", "act_Source",
     "FICTIONAL: Federal Register notice 92 FR 00000; HTSUS heading 9903.XX.01", "TIM Input",
     "Brief Part 4. Citation of the notice, order or decree."),
    ("Effective date", "act_EffDate", dt.date(2027, 3, 1), "TIM Input Date",
     "Brief Part 5. Shown on the Summary; not applied by the formulas."),
    ("Effective time and basis", "act_Basis",
     "12:01 a.m. ET; goods entered for consumption, or withdrawn from warehouse for consumption, on or after the effective date",
     "TIM Input", "Brief Part 5. Quote the notice wording."),
    ("In-transit provision", "act_InTransit",
     "FICTIONAL: not covered if loaded on the final mode of transit before the effective date and entered before 2027-04-30",
     "TIM Input", "Brief Part 5. Not applied by the formulas; Standard Cost handles in-transit stock separately."),
    ("End date (blank if none)", "act_EndDate", None, "TIM Input Date",
     "Brief Part 5. Not applied by the formulas."),
    ("Default additional rate (% of customs value)", "act_Rate", 0.25, "TIM Input Pct",
     "Brief Part 7. Applies to every in-scope item that matches no rate line (columns Y to AI). "
     "Leave blank if every in-scope item must match a rate line. Ad valorem only."),
    ("Rate change type (ADD or REPLACE)", "act_RateType", "ADD", "TIM Input Text",
     "ADD: the rates are added to the item's current rate. REPLACE: they replace this action's earlier rate "
     "(B17, or column AD of a rate line)."),
    ("Earlier rate of this action being replaced (default)", "act_PriorRate", None, "TIM Input Pct",
     "REPLACE only. Brief Part 7. Must already be inside each item's current rate on Items."),
    ("USMCA-qualifying goods (Y, PARTIAL, N or N/A)", "act_USMCAExempt", "N/A", "TIM Input Text",
     "Brief 6.5. Y = exempt: items flagged USMCA-qualifying on Items are out of scope. PARTIAL = they stay in scope and "
     "get relief through a rate line with USMCA = Y. N = no relief. N/A = no USMCA-party origin in scope."),
    ("Stacking interpretation", "act_Stacking",
     "Adds to the HTSUS column 1 General rate. FICTIONAL 9903.YY.01 (an earlier 10% on listed goods of Country X) does not "
     "apply to goods that pay 9903.XX.01 (S1 clause 4); SAMPLE-003 pays it today (determination, column S).",
     "TIM Input", "Brief 6.6, quoted. Duties the action switches off go in column AF (per rate line) or column S (per item)."),
    ("Duty base for the default rate (FULL or PART)", "act_DefaultBase", "FULL", "TIM Input Text",
     "Brief Part 7. FULL = the whole customs value. PART = only part of it (for example metal or non-US content): "
     "Standard Cost enters each item's share on Items column M, from the source Research names."),
    ("Research maker", "act_Maker", "Hire A", "TIM Input", "Brief Part 12."),
    ("Research checker", "act_Checker", "Senior", "TIM Input", "Brief Part 12."),
    ("Date brief checked", "act_Checked", dt.date(2027, 2, 18), "TIM Input Date", "Brief Part 12."),
    ("Model version", "act_ModelVersion", "v1.1 (SAMPLE)", "TIM Input Text",
     "Matches the latest line on the Review tab and the file name."),
]
FIELD_ROW = {d: FIRST + i for i, (_, d, *_rest) in enumerate(ACTION_FIELDS)}
assert FIELD_ROW["act_Authority"] == 8 and FIELD_ROW["act_Source"] == 10 and FIELD_ROW["act_Stacking"] == 19
assert FIELD_ROW["act_ModelVersion"] == 24

AUTHORITIES = ["Section 232", "Section 301", "Section 201", "Section 122", "Section 338", "IEEPA", "AD/CVD",
               "Canada surtax order", "Mexico decree (DOF)", "Other"]

# Action list blocks: (first column, last column). Headers on row 4.
SCOPE_COLS = ("E", "H")   # E prefix, F digits (formula), G description, H rate group
EXCL_COLS = ("I", "K")    # I prefix, J digits (formula), K basis
ORIGIN_COLS = ("M", "N")  # M code, N note
DET_COLS = ("P", "W")     # P item, Q import country, R IN/OUT, S displaced, T reason, U ref, V key (f), W check (f)
RATE_COLS = ("Y", "AI")   # Y ID, Z group, AA origin, AB USMCA, AC rate, AD earlier, AE base, AF displaced, AG ref,
#                           AH items using (f), AI check (f)

SAMPLE_SCOPE_HTS = [("7318", "FICTIONAL Annex I: heading 7318, screws, bolts, nuts, washers, rivets and similar, of iron or steel",
                     "9903.XX.01")]
SAMPLE_EXCL_HTS = [("7318.24", "FICTIONAL Annex II (1): cotters and cotter pins")]
SAMPLE_COO = [("XX", "Country X (FICTIONAL; XX is an ISO user-assigned code)")]
SAMPLE_DETERMINATIONS = [
    ("SAMPLE-005", "US", "OUT", None,
     "FICTIONAL Annex II (2): the supplier certifies every shipment for civil aircraft; the certificate of conformity "
     "is filed with each entry", "Brief 6.4"),
    ("SAMPLE-003", "US", None, 0.10,
     "Pays FICTIONAL 9903.YY.01 (10%) today; S1 clause 4 switches it off once 9903.XX.01 applies", "Brief 6.6, Q5"),
]
SAMPLE_RATE_LINES = []  # the fictional action has one rate (B15); the guide's drills add rate lines


def build_action(wb, calc_col):
    ws = wb.create_sheet("Action")
    title(ws, "Action: the WHAT (Research fills this tab from the checked Tariff-Impact Brief)",
          "Standard Cost reads this tab but never changes it. Questions go to Research on the Jira story.",
          "SAMPLE: the values below describe a FICTIONAL action for training. No such action exists.")
    headers(ws, 4, 1, ["Field", "Value", "Where it comes from"])
    for i, (label, dname, value, style, note) in enumerate(ACTION_FIELDS):
        r = FIRST + i
        put(ws, f"A{r}", label, "TIM Label")
        put(ws, f"B{r}", value, style)
        put(ws, f"C{r}", note, "TIM Note")
        name(wb, dname, f"Action!$B${r}")

    digits = '=IF(TRIM({c}{r}&"")="","",SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({c}{r}&""),".","")," ",""),"-",""))'

    # HTS scope list and exclusion list.
    headers(ws, 4, 5, ["In-scope HTS prefix (as written)", "Digits (formula)", "Annex line or description (source)",
                       "Rate group (blank unless rate lines use it)"])
    headers(ws, 4, 9, ["Excluded HTS prefix (as written)", "Digits (formula)", "Exclusion basis (source)"])
    for r in range(FIRST, HTS_LAST + 1):
        put(ws, f"E{r}", None, "TIM Input Text")
        put(ws, f"F{r}", digits.format(c="E", r=r), "TIM Calc")
        put(ws, f"G{r}", None, "TIM Input")
        put(ws, f"H{r}", None, "TIM Input Text")
        put(ws, f"I{r}", None, "TIM Input Text")
        put(ws, f"J{r}", digits.format(c="I", r=r), "TIM Calc")
        put(ws, f"K{r}", None, "TIM Input")
    for i, (prefix, desc, group) in enumerate(SAMPLE_SCOPE_HTS):
        r = FIRST + i
        ws[f"E{r}"].value, ws[f"G{r}"].value, ws[f"H{r}"].value = prefix, desc, group
    for i, (prefix, desc) in enumerate(SAMPLE_EXCL_HTS):
        ws[f"I{FIRST + i}"].value, ws[f"K{FIRST + i}"].value = prefix, desc

    # Origin list.
    headers(ws, 4, 13, ["Origin in scope (ISO 2)", "Country or note"])
    for r in range(FIRST, ORIGIN_LAST + 1):
        put(ws, f"M{r}", None, "TIM Input Text")
        put(ws, f"N{r}", None, "TIM Input")
    for i, (code, desc) in enumerate(SAMPLE_COO):
        ws[f"M{FIRST + i}"].value, ws[f"N{FIRST + i}"].value = code, desc

    # Item-level determinations, keyed by item AND import country.
    key_rng = f"Calc!${calc_col['key']}${FIRST}:${calc_col['key']}${LAST}"
    headers(ws, 4, 16, ["Item number", "Import country", "Determination (IN, OUT or blank = scope rules)",
                        "Other duty this action displaces on the item (% of customs value)", "Reason", "Brief ref",
                        "Key (formula)", "Check (formula)"])
    for r in range(FIRST, DET_LAST + 1):
        put(ws, f"P{r}", None, "TIM Input Text")
        put(ws, f"Q{r}", None, "TIM Input Text")
        put(ws, f"R{r}", None, "TIM Input Text")
        put(ws, f"S{r}", None, "TIM Input Pct")
        put(ws, f"T{r}", None, "TIM Input")
        put(ws, f"U{r}", None, "TIM Input")
        put(ws, f"V{r}", f'=IF(TRIM(P{r}&"")="","",UPPER(TRIM(P{r}&""))&"/"&UPPER(TRIM(Q{r}&"")))', "TIM Calc")
        put(ws, f"W{r}",
            f'=IF(V{r}="","",IF(UPPER(TRIM(Q{r}&""))<>UPPER(TRIM(act_Country&"")),"CHECK: import country must match B7",'
            f'IF(COUNTIF(det_Key,V{r})>1,"CHECK: listed twice",'
            f'IF(COUNTIF({key_rng},V{r})=0,"CHECK: no Items row for this item and import country",'
            f'IF(AND(TRIM(R{r}&"")="",NOT(ISNUMBER(S{r}))),"CHECK: enter IN, OUT or a displaced rate","OK")))))',
            "TIM Calc")
    for i, (item, ctry, det, displaced, reason, ref) in enumerate(SAMPLE_DETERMINATIONS):
        r = FIRST + i
        for col, v in zip("PQRSTU", (item, ctry, det, displaced, reason, ref)):
            ws[f"{col}{r}"].value = v

    # Rate lines.
    line_rng = f"Calc!${calc_col['line']}${FIRST}:${calc_col['line']}${LAST}"
    in_rng = f"Calc!${calc_col['inscope']}${FIRST}:${calc_col['inscope']}${LAST}"
    headers(ws, 4, 25, ["Rate line ID (brief Part 7)", "Rate group (blank = all)", "Origin (blank = all listed)",
                        "USMCA-qualifying (Y, N or blank = either)", "Additional rate (% of customs value)",
                        "Earlier rate replaced (REPLACE only; blank = B17)", "Duty base (FULL or PART; blank = FULL)",
                        "Other duty this line displaces (% of customs value)", "Brief ref",
                        "In-scope items using this line (formula)", "Check (formula)"])
    for r in range(FIRST, RATE_LAST + 1):
        for col, style in (("Y", "TIM Input Text"), ("Z", "TIM Input Text"), ("AA", "TIM Input Text"),
                           ("AB", "TIM Input Text"), ("AC", "TIM Input Pct"), ("AD", "TIM Input Pct"),
                           ("AE", "TIM Input Text"), ("AF", "TIM Input Pct"), ("AG", "TIM Input")):
            put(ws, f"{col}{r}", None, style)
        put(ws, f"AH{r}", f'=IF(TRIM(Y{r}&"")="","",COUNTIFS({line_rng},TRIM(Y{r}&""),{in_rng},"Y"))', "TIM Calc Whole")
        put(ws, f"AI{r}",
            f'=IF(AND(TRIM(Y{r}&"")="",TRIM(Z{r}&"")="",TRIM(AA{r}&"")="",TRIM(AB{r}&"")="",AC{r}&""=""),"",'
            f'IF(TRIM(Y{r}&"")="","CHECK: line ID missing",IF(NOT(ISNUMBER(AC{r})),"CHECK: rate missing",'
            f'IF(COUNTIF(rate_ID,TRIM(Y{r}&""))>1,"CHECK: line ID listed twice",'
            f'IF(AND(TRIM(Z{r}&"")<>"",COUNTIF(scope_Group,TRIM(Z{r}&""))=0),"CHECK: rate group not on the scope list",'
            f'IF(AND(TRIM(AA{r}&"")<>"",COUNTIF(scope_COO,TRIM(AA{r}&""))=0),"CHECK: origin not on the origin list",'
            f'"OK"))))))', "TIM Calc")
    for i, row in enumerate(SAMPLE_RATE_LINES):
        for col, v in zip(["Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG"], row):
            ws[f"{col}{FIRST + i}"].value = v

    name(wb, "scope_HTS", f"Action!$F${FIRST}:$F${HTS_LAST}")
    name(wb, "scope_Group", f"Action!$H${FIRST}:$H${HTS_LAST}")
    name(wb, "excl_HTS", f"Action!$J${FIRST}:$J${HTS_LAST}")
    name(wb, "scope_COO", f"Action!$M${FIRST}:$M${ORIGIN_LAST}")
    name(wb, "det_Status", f"Action!$R${FIRST}:$R${DET_LAST}")
    name(wb, "det_Displaced", f"Action!$S${FIRST}:$S${DET_LAST}")
    name(wb, "det_Key", f"Action!$V${FIRST}:$V${DET_LAST}")
    name(wb, "det_Check", f"Action!$W${FIRST}:$W${DET_LAST}")
    for label, col in (("rate_ID", "Y"), ("rate_Group", "Z"), ("rate_COO", "AA"), ("rate_USMCA", "AB"),
                       ("rate_Rate", "AC"), ("rate_Prior", "AD"), ("rate_Base", "AE"), ("rate_Displaced", "AF"),
                       ("rate_Check", "AI")):
        name(wb, label, f"Action!${col}${FIRST}:${col}${RATE_LAST}")

    list_dv(ws, ["US", "CA", "MX"], f"B{FIELD_ROW['act_Country']}")
    list_dv(ws, AUTHORITIES, f"B{FIELD_ROW['act_Authority']}")
    list_dv(ws, ["ADD", "REPLACE"], f"B{FIELD_ROW['act_RateType']}",
            "ADD = on top of the current rate. REPLACE = replaces this action's earlier rate.")
    list_dv(ws, ["Y", "PARTIAL", "N", "N/A"], f"B{FIELD_ROW['act_USMCAExempt']}",
            "Y = exempt. PARTIAL = relief through a rate line with USMCA = Y. N = no relief.")
    list_dv(ws, ["FULL", "PART"], f"B{FIELD_ROW['act_DefaultBase']}")
    for key in ("act_Rate", "act_PriorRate"):
        number_dv(ws, "decimal", f"B{FIELD_ROW[key]}", 0, 10, "Enter a rate between 0% and 1000%, as a percentage.")
    for key in ("act_EffDate", "act_EndDate", "act_Checked"):
        number_dv(ws, "date", f"B{FIELD_ROW[key]}", 36526, 73051, "Enter a date (yyyy-mm-dd).")
    length_dv(ws, f"M{FIRST}:M{ORIGIN_LAST}", 2, "Use the 2-letter ISO country code, no spaces.")
    list_dv(ws, ["US", "CA", "MX"], f"Q{FIRST}:Q{DET_LAST}")
    list_dv(ws, ["IN", "OUT"], f"R{FIRST}:R{DET_LAST}", "Blank = the scope rules decide (use it to record a displaced duty only).")
    number_dv(ws, "decimal", f"S{FIRST}:S{DET_LAST}", 0, 10, "Enter a rate as a percentage, for example 10%.")
    length_dv(ws, f"AA{FIRST}:AA{RATE_LAST}", 2, "Use the 2-letter ISO country code, no spaces.")
    list_dv(ws, ["Y", "N"], f"AB{FIRST}:AB{RATE_LAST}", "Y = only USMCA-qualifying items; N = only non-qualifying; blank = either.")
    for col in ("AC", "AD", "AF"):
        number_dv(ws, "decimal", f"{col}{FIRST}:{col}{RATE_LAST}", 0, 10, "Enter a rate as a percentage, for example 25%.")
    list_dv(ws, ["FULL", "PART"], f"AE{FIRST}:AE{RATE_LAST}")

    ws.conditional_formatting.add(f"W{FIRST}:W{DET_LAST}", FormulaRule(formula=[f'LEFT($W{FIRST},5)="CHECK"'], fill=RED))
    ws.conditional_formatting.add(f"AI{FIRST}:AI{RATE_LAST}", FormulaRule(formula=[f'LEFT($AI{FIRST},5)="CHECK"'], fill=RED))

    widths(ws, {"A": 30, "B": 44, "C": 40, "D": 2, "E": 16, "F": 11, "G": 36, "H": 14, "I": 16, "J": 11, "K": 32,
                "L": 2, "M": 9, "N": 28, "O": 2, "P": 14, "Q": 9, "R": 14, "S": 14, "T": 36, "U": 12, "V": 16, "W": 30,
                "X": 2, "Y": 10, "Z": 13, "AA": 9, "AB": 12, "AC": 11, "AD": 12, "AE": 11, "AF": 13, "AG": 12,
                "AH": 11, "AI": 30})
    ws.row_dimensions[4].height = 58
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = "2E75B6"
    protect(ws)


# ------------------------------------------------------------------- Items ----

ITEM_HEADERS = ["Item number", "Description", "Import country (US, CA, MX)", "HTS, 10 digits (text)",
                "Country of origin (ISO 2)", "USMCA qualifying (Y/N)", "Unit of measure (Oracle primary UOM)",
                "Unit customs value (per unit in G)", "Quantity for the period (in unit G)",
                "Current duty rate (all duties now paid, % of customs value)", "Current standard cost (per unit in G)",
                "Sell price (per unit in G; blank if not sold)",
                "Dutiable share of customs value (PART base only; blank = 100%)", "Source and notes"]
ITEM_STYLES = ["TIM Input Text", "TIM Input", "TIM Input Text", "TIM Input Text", "TIM Input Text", "TIM Input Text",
               "TIM Input Text", "TIM Input Money", "TIM Input Qty", "TIM Input Pct", "TIM Input Money", "TIM Input Money",
               "TIM Input Pct", "TIM Input"]

# SYNTHETIC rows. The first 6 HTS digits are real HS subheadings; the last 4 digits are invented.
# Values, rates, costs and prices are invented. XX and XY are ISO user-assigned codes for fictional countries.
# (item, description, import country, HTS, origin, USMCA, UOM, customs value, quantity, current rate, std cost,
#  sell price, dutiable share, notes)
SAMPLE_ITEMS = [
    ("SAMPLE-001", "Hex cap screw, steel, zinc plated", "US", "7318.15.9001", "XX", "N", "EA", 4.20, 120000, 0.05, 4.95, 7.40,
     None, "In scope by rules"),
    ("SAMPLE-002", "Hex nut, steel", "US", "7318169001", "XX", "N", "EA", 0.85, 400000, 0.00, 1.02, 1.55, None,
     "In scope; HTS pasted without dots"),
    ("SAMPLE-003", "Flat washer, steel", "US", "7318.22.9001", "XX", "N", "EA", 0.12, 900000, 0.10, 0.16, 0.24, None,
     "In scope; current rate includes FICTIONAL 9903.YY.01 10%, which the action displaces"),
    ("SAMPLE-004", "Cotter pin, steel", "US", "7318.24.9001", "XX", "N", "EA", 0.30, 50000, 0.00, 0.36, 0.60, None,
     "Out: excluded subheading"),
    ("SAMPLE-005", "Bolt certified for civil aircraft (certificate with every entry)", "US", "7318.15.9002", "XX", "N", "EA",
     12.50, 8000, 0.025, 14.10, 22.00, None, "Out: Research determination"),
    ("SAMPLE-006", "Hex cap screw, steel, from Country Y", "US", "7318.15.9001", "XY", "N", "EA", 3.90, 60000, 0.05, 4.60, 7.40,
     None, "Out: origin not listed"),
    ("SAMPLE-007", "Carriage bolt, steel, from Mexico, USMCA", "US", "7318.15.9003", "MX", "Y", "EA", 2.10, 150000, 0.00, 2.45,
     3.90, None, "Out: origin not listed"),
    ("SAMPLE-008", "Carriage bolt, steel, from Mexico, non-qualifying", "US", "7318.15.9003", "MX", "N", "EA", 2.30, 40000, 0.05,
     2.70, 3.90, None, "Out: origin not listed"),
    ("SAMPLE-009", "Self-tapping screw, steel, from Canada", "US", "7318.14.9001", "CA", "N", "EA", 0.40, 300000, 0.05, 0.47,
     0.80, None, "Out: origin not listed"),
    ("SAMPLE-010", "Mounting bracket kit, steel, with fasteners", "US", "7326.90.9001", "XX", "N", "EA", 6.00, 20000, 0.029,
     6.80, 11.50, None, "Out: HTS not listed"),
    ("SAMPLE-011", "Ball valve, brass", "US", "8481.80.9001", "XX", "N", "EA", 18.00, 5000, 0.04, 20.50, 34.00, None,
     "Out: HTS not listed"),
    ("SAMPLE-012", "Rivet, steel", "US", "7318.23.9001", "XX", "N", "EA", 0.05, 0, 0.00, 0.06, 0.10, None,
     "In scope; no forecast quantity"),
    ("SAMPLE-013", "Hex cap screw, steel, imported into Canada", "CA", "7318.15.9001", "XX", "N", "EA", 4.20, 30000, 0.00, 4.95,
     7.40, None, "Out: different import country"),
    ("SAMPLE-014", "Machine screw, steel, for internal use (not sold)", "US", "7318.15.9004", "XX", "N", "EA", 0.50, 10000,
     0.05, 0.60, None, None, "In scope; no sell price"),
]
assert len(ITEM_HEADERS) == len(ITEM_STYLES) == len(SAMPLE_ITEMS[0])

SETTINGS = [  # (label cell, label, value cell, defined name, sample value, note)
    ("J1", "Quantity period (label)", "K1", "qty_Period", "Annual",
     "Period the quantities in column I cover, for example Annual or Q2 2027. Shown on the Summary."),
    ("J2", "Currency of all values", "K2", "mdl_Currency", "USD",
     "[COMPANY: model currency]. Every value, cost and price on this tab is in this currency."),
]


def build_items(wb):
    ws = wb.create_sheet("Items")
    title(ws, "Items: the $ inputs (Standard Cost fills this tab from the item extract)",
          "Source: [COMPANY: Oracle cost extract or warehouse query]. One row per item and import country. "
          "Values per Oracle primary UOM. Paste values only.",
          f"SAMPLE ROWS {FIRST}-{FIRST + len(SAMPLE_ITEMS) - 1} ARE SYNTHETIC: invented items, values and rates; "
          "HTS last 4 digits are made up. Clear before real use.")
    for label_cell, label, value_cell, dname, value, note in SETTINGS:
        put(ws, label_cell, label, "TIM Label")
        put(ws, value_cell, value, "TIM Input Text")
        put(ws, value_cell.replace("K", "L"), note, "TIM Note")
        name(wb, dname, f"Items!${value_cell[0]}${value_cell[1:]}")
    headers(ws, 4, 1, ITEM_HEADERS)
    for r in range(FIRST, LAST + 1):
        for c, style in enumerate(ITEM_STYLES, 1):
            put(ws, f"{get_column_letter(c)}{r}", None, style)
    for i, row in enumerate(SAMPLE_ITEMS):
        for c, v in enumerate(row, 1):
            ws.cell(row=FIRST + i, column=c).value = v
    rng = lambda col: f"{col}{FIRST}:{col}{LAST}"  # noqa: E731
    list_dv(ws, ["US", "CA", "MX"], rng("C"))
    length_dv(ws, rng("E"), 2, "Use the 2-letter ISO country code, no spaces.")
    list_dv(ws, ["Y", "N"], rng("F"), "Y only if a valid USMCA certification supports the claim.")
    for col in "HKL":
        number_dv(ws, "decimal", rng(col), 0, 100000000, "Enter a number of 0 or more.")
    number_dv(ws, "decimal", rng("I"), 0, 2000000000, "Enter a quantity of 0 or more, in the unit in column G.")
    number_dv(ws, "decimal", rng("J"), 0, 10, "Enter the rate as a percentage, for example 2.5%.")
    number_dv(ws, "decimal", rng("M"), 0, 1, "Enter a share between 0% and 100%.")
    widths(ws, {"A": 14, "B": 38, "C": 10, "D": 15, "E": 10, "F": 10, "G": 11, "H": 12, "I": 13, "J": 15, "K": 13,
                "L": 13, "M": 14, "N": 34})
    ws.row_dimensions[4].height = 58
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:N{LAST}"
    ws.sheet_properties.tabColor = "C55A11"
    protect(ws, sort=True, filter_=True)


# -------------------------------------------------------------------- Calc ----

# (key, header, formula template, style, width). Templates use {r} for the row and {key} for another column's letter.
LINE_IDX = "MOD(${r_score}{r},100)"
RATE_MATCH = ('ISNUMBER(rate_Rate)*((TRIM(rate_Group&"")="")+(TRIM(rate_Group&"")=${group}{r})>0)'
              '*((rate_COO="")+(rate_COO=${coo}{r})>0)'
              '*((rate_USMCA="")+(rate_USMCA=IF(UPPER(TRIM(Items!F{r}&""))="Y","Y","N"))>0)')
RATE_SPEC = '(1+(TRIM(rate_Group&"")<>"")+2*(rate_USMCA<>"")+4*(rate_COO<>""))'
SCOPE_MATCH = '(scope_HTS<>"")*(LEFT(${hts}{r},LEN(scope_HTS))=scope_HTS)'

CALC_COLUMNS = [
    ("item", "Item number", '=IF(TRIM(Items!A{r}&"")="","",TRIM(Items!A{r}&""))', "TIM Calc", 14),
    ("desc", "Description", '=IF(${item}{r}="","",Items!B{r}&"")', "TIM Calc", 28),
    ("hts", "HTS digits",
     '=IF(${item}{r}="","",SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM(Items!D{r}&""),".","")," ",""),"-",""))', "TIM Calc", 12),
    ("chap", "HTS chapter", '=IF(${item}{r}="","",LEFT(${hts}{r},2))', "TIM Calc", 8),
    ("coo", "Origin", '=IF(${item}{r}="","",UPPER(TRIM(Items!E{r}&"")))', "TIM Calc", 8),
    ("ctry_ok", "Import country matches (Y/N)",
     '=IF(${item}{r}="","",IF(UPPER(TRIM(Items!C{r}&""))=UPPER(TRIM(act_Country&"")),"Y","N"))', "TIM Calc", 10),
    ("on_scope", "HTS on scope list (Y/N)", '=IF(${item}{r}="","",IF(${s_score}{r}>0,"Y","N"))', "TIM Calc", 10),
    ("on_excl", "HTS on exclusion list (Y/N)",
     '=IF(${item}{r}="","",IF(SUMPRODUCT((excl_HTS<>"")*(LEFT(${hts}{r},LEN(excl_HTS))=excl_HTS))>0,"Y","N"))', "TIM Calc", 10),
    ("coo_ok", "Origin on scope list (Y/N)",
     '=IF(${item}{r}="","",IF(${coo}{r}="","N",IF(COUNTIF(scope_COO,${coo}{r})>0,"Y","N")))', "TIM Calc", 10),
    ("usmca_ex", "USMCA exempt (Y/N)",
     '=IF(${item}{r}="","",IF(AND(UPPER(TRIM(act_USMCAExempt&""))="Y",UPPER(TRIM(Items!F{r}&""))="Y"),"Y","N"))',
     "TIM Calc", 10),
    ("det", "Research determination",
     '=IF(${item}{r}="","",IF(${det_row}{r}=0,"",UPPER(TRIM(INDEX(det_Status,${det_row}{r})&""))))', "TIM Calc", 12),
    ("inscope", "In scope (Y/N)",
     '=IF(${item}{r}="","",IF(${ctry_ok}{r}="N","N",IF(${det}{r}="IN","Y",IF(${det}{r}="OUT","N",'
     'IF(AND(${on_scope}{r}="Y",${on_excl}{r}="N",${coo_ok}{r}="Y",${usmca_ex}{r}="N"),"Y","N")))))', "TIM Calc", 9),
    ("reason", "Scope reason",
     '=IF(${item}{r}="","",IF(${ctry_ok}{r}="N","Out: different import country",'
     'IF(${det}{r}="IN","In: Research determination",IF(${det}{r}="OUT","Out: Research determination",'
     'IF(${on_scope}{r}="N","Out: HTS not on scope list",IF(${on_excl}{r}="Y","Out: HTS on exclusion list",'
     'IF(${coo_ok}{r}="N","Out: origin not on scope list",'
     'IF(${usmca_ex}{r}="Y","Out: USMCA-qualifying, exempt","In: matches scope rules"))))))))', "TIM Calc", 28),
    ("group", "Rate group (from scope list)",
     '=IF(OR(${item}{r}="",${s_score}{r}=0),"",TRIM(INDEX(scope_Group,MOD(${s_score}{r},1000))&""))', "TIM Calc", 12),
    ("line", "Rate line used",
     '=IF(${inscope}{r}<>"Y","",IF(${r_score}{r}=0,"Default (B15)",TRIM(INDEX(rate_ID,' + LINE_IDX + ')&"")))',
     "TIM Calc", 12),
    ("base", "Duty base (FULL or PART)",
     '=IF(${inscope}{r}<>"Y","",IF(${r_score}{r}=0,IF(UPPER(TRIM(act_DefaultBase&""))="PART","PART","FULL"),'
     'IF(UPPER(TRIM(INDEX(rate_Base,' + LINE_IDX + ')&""))="PART","PART","FULL")))', "TIM Calc", 9),
    ("cur", "Current duty rate", '=IF(${item}{r}="","",N(Items!J{r}))', "TIM Calc Pct", 10),
    ("add", "Additional rate (this action)",
     '=IF(${item}{r}="","",IF(${inscope}{r}<>"Y",0,IF(${r_score}{r}=0,N(act_Rate),N(INDEX(rate_Rate,' + LINE_IDX + ')))))',
     "TIM Calc Pct", 10),
    ("prior", "Earlier rate replaced",
     '=IF(${item}{r}="","",IF(OR(${inscope}{r}<>"Y",UPPER(TRIM(act_RateType&""))<>"REPLACE"),0,'
     'IF(${r_score}{r}=0,N(act_PriorRate),IF(ISNUMBER(INDEX(rate_Prior,' + LINE_IDX + ')),'
     'INDEX(rate_Prior,' + LINE_IDX + '),N(act_PriorRate)))))', "TIM Calc Pct", 10),
    ("share", "Dutiable share applied",
     '=IF(${item}{r}="","",IF(AND(${inscope}{r}="Y",${base}{r}="PART",ISNUMBER(Items!M{r})),Items!M{r},1))',
     "TIM Calc Pct", 10),
    ("displ", "Other duty displaced",
     '=IF(${item}{r}="","",IF(${inscope}{r}<>"Y",0,IF(${det_row}{r}>0,'
     'IF(ISNUMBER(INDEX(det_Displaced,${det_row}{r})),INDEX(det_Displaced,${det_row}{r}),'
     'IF(${r_score}{r}=0,0,N(INDEX(rate_Displaced,' + LINE_IDX + ')))),'
     'IF(${r_score}{r}=0,0,N(INDEX(rate_Displaced,' + LINE_IDX + '))))))', "TIM Calc Pct", 10),
    ("new", "New duty rate (all duties, % of full customs value)",
     '=IF(${item}{r}="","",IF(${inscope}{r}="Y",${cur}{r}+${share}{r}*(${add}{r}-${prior}{r})-${displ}{r},${cur}{r}))',
     "TIM Calc Pct", 11),
    ("dut_b", "Duty per unit before", '=IF(${item}{r}="","",N(Items!H{r})*${cur}{r})', "TIM Calc Unit", 11),
    ("dut_a", "Duty per unit after", '=IF(${item}{r}="","",N(Items!H{r})*${new}{r})', "TIM Calc Unit", 11),
    ("chg", "Landed cost change per unit", '=IF(${item}{r}="","",${dut_a}{r}-${dut_b}{r})', "TIM Calc Unit", 11),
    ("newstd", "Est. new standard cost per unit", '=IF(${item}{r}="","",N(Items!K{r})+${chg}{r})', "TIM Calc Unit", 12),
    ("m_b", "Margin % before",
     '=IF(${item}{r}="","",IF(N(Items!L{r})=0,"",(N(Items!L{r})-N(Items!K{r}))/N(Items!L{r})))', "TIM Calc Pct", 10),
    ("m_a", "Margin % after",
     '=IF(${item}{r}="","",IF(N(Items!L{r})=0,"",(N(Items!L{r})-${newstd}{r})/N(Items!L{r})))', "TIM Calc Pct", 10),
    ("m_chg", "Margin change (points)", '=IF(OR(${m_b}{r}="",${m_a}{r}=""),"",${m_a}{r}-${m_b}{r})', "TIM Calc Pct", 10),
    ("qty", "Quantity for the period", '=IF(${item}{r}="","",N(Items!I{r}))', "TIM Calc Whole", 11),
    ("impact", "Impact for the period (change x quantity)", '=IF(${item}{r}="","",${chg}{r}*${qty}{r})', "TIM Calc Money", 13),
    ("priced", "Has a sell price (Y/N)", '=IF(${item}{r}="","",IF(N(Items!L{r})>0,"Y","N"))', "TIM Calc", 9),
    ("rev", "Revenue for the period (items with a sell price)",
     '=IF(${item}{r}="","",IF(${priced}{r}="Y",N(Items!L{r})*${qty}{r},0))', "TIM Calc Money", 13),
    ("mg_b", "Margin $ before (items with a sell price)",
     '=IF(${item}{r}="","",IF(${priced}{r}="Y",(N(Items!L{r})-N(Items!K{r}))*${qty}{r},0))', "TIM Calc Money", 13),
    ("mg_a", "Margin $ after (items with a sell price)",
     '=IF(${item}{r}="","",IF(${priced}{r}="Y",${mg_b}{r}-${impact}{r},0))', "TIM Calc Money", 13),
    ("check", "Data check",
     '=IF(${item}{r}="","",'
     'IF(COUNTIF(${key}$' + str(FIRST) + ':${key}$' + str(LAST) + ',${key}{r})>1,"CHECK: duplicate item and import country",'
     'IF(TRIM(Items!C{r}&"")="","CHECK: import country missing",'
     'IF(OR(LEN(${hts}{r})<>10,ISERROR(VALUE(${hts}{r}))),"CHECK: HTS must be 10 digits",'
     'IF(LEN(${coo}{r})<>2,"CHECK: origin code missing or not 2 letters",'
     'IF(TRIM(Items!G{r}&"")="","CHECK: unit of measure missing",'
     'IF(OR(Items!H{r}="",Items!J{r}="",Items!K{r}=""),"CHECK: value, rate or cost missing",'
     'IF(OR(NOT(ISNUMBER(Items!H{r})),NOT(ISNUMBER(Items!J{r})),NOT(ISNUMBER(Items!K{r})),'
     'AND(Items!I{r}<>"",NOT(ISNUMBER(Items!I{r}))),AND(Items!L{r}<>"",NOT(ISNUMBER(Items!L{r}))),'
     'AND(Items!M{r}<>"",NOT(ISNUMBER(Items!M{r})))),"CHECK: text in a number column",'
     'IF(AND(UPPER(TRIM(Items!F{r}&""))="Y",${coo}{r}<>"US",${coo}{r}<>"CA",${coo}{r}<>"MX"),'
     '"CHECK: USMCA flag on an origin other than US, CA or MX",'
     'IF(AND(${inscope}{r}="Y",${r_score}{r}=0,NOT(ISNUMBER(act_Rate))),"CHECK: in scope but no rate (add a rate line or B15)",'
     'IF(AND(${inscope}{r}="Y",${r_ties}{r}>1),"CHECK: two rate lines match equally",'
     'IF(AND(${inscope}{r}="Y",${s_ties}{r}>1),"CHECK: HTS prefix listed twice on the scope list",'
     'IF(AND(${inscope}{r}="Y",${base}{r}="PART",NOT(ISNUMBER(Items!M{r}))),"CHECK: dutiable share missing (PART base)",'
     'IF(AND(${inscope}{r}="Y",${base}{r}="FULL",ISNUMBER(Items!M{r}),N(Items!M{r})<>1),'
     '"CHECK: dutiable share entered, but the rate taxes the full value",'
     'IF(AND(${inscope}{r}="Y",${cur}{r}+0.0000001<${share}{r}*${prior}{r}+${displ}{r}),'
     '"CHECK: current rate is below the duty being replaced or displaced",'
     'IF(AND(${inscope}{r}="Y",${qty}{r}=0),"NOTE: in scope, no quantity",'
     'IF(AND(${inscope}{r}="Y",${priced}{r}="N"),"NOTE: in scope, no sell price (left out of margin %)",'
     'IF(AND(${inscope}{r}="Y",N(Items!H{r})>N(Items!K{r})),"NOTE: customs value above standard cost (check unit of measure)",'
     '"OK"))))))))))))))))))', "TIM Calc", 30),
    ("rank", "Rank key (helper for Summary top 10)",
     '=IF(${item}{r}="","",IF(AND(${inscope}{r}="Y",${impact}{r}<>0),ABS(${impact}{r})+ROW()/1000000000,""))',
     "TIM Calc Key", 12),
    ("key", "Item and import country key (helper)",
     '=IF(${item}{r}="","",${item}{r}&"/"&UPPER(TRIM(Items!C{r}&"")))', "TIM Calc", 16),
    ("s_score", "Scope match score (helper)",
     '=IF(${item}{r}="",0,SUMPRODUCT(MAX(' + SCOPE_MATCH + '*(LEN(scope_HTS)*1000+ROW(scope_HTS)-' + str(FIRST - 1) + '))))',
     "TIM Calc Whole", 10),
    ("s_ties", "Scope prefixes at best length (helper)",
     '=IF(${s_score}{r}=0,0,SUMPRODUCT(' + SCOPE_MATCH + '*(LEN(scope_HTS)=INT(${s_score}{r}/1000))))', "TIM Calc Whole", 10),
    ("r_score", "Rate line score (helper)",
     '=IF(${inscope}{r}<>"Y",0,SUMPRODUCT(MAX(' + RATE_MATCH + '*(' + RATE_SPEC + '*100+ROW(rate_Rate)-'
     + str(FIRST - 1) + '))))', "TIM Calc Whole", 10),
    ("r_ties", "Rate lines at best score (helper)",
     '=IF(${r_score}{r}=0,0,SUMPRODUCT(' + RATE_MATCH + '*(' + RATE_SPEC + '=INT(${r_score}{r}/100))))', "TIM Calc Whole", 10),
    ("det_row", "Determination row (helper)",
     '=IF(${item}{r}="",0,IFERROR(MATCH(${key}{r},det_Key,0),0))', "TIM Calc Whole", 10),
]
CALC_COL = {key: get_column_letter(i + 1) for i, (key, *_rest) in enumerate(CALC_COLUMNS)}
assert len(CALC_COL) == len(CALC_COLUMNS)


def calc_formula(template, r):
    return template.format(r=r, **CALC_COL)


def build_calc(wb):
    ws = wb.create_sheet("Calc")
    title(ws, "Calc: formulas only (locked; nothing to type on this tab)",
          "Row n here is row n on Items. Grey = formula. Scope and rates come from the Action tab; values come from the Items tab.")
    headers(ws, 4, 1, [h for _, h, *_ in CALC_COLUMNS])
    for c, (_, _, template, style, width) in enumerate(CALC_COLUMNS, 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = width
        for r in range(FIRST, LAST + 1):
            put(ws, f"{letter}{r}", calc_formula(template, r), style)
    ins, chk = CALC_COL["inscope"], CALC_COL["check"]
    ws.conditional_formatting.add(f"{ins}{FIRST}:{ins}{LAST}", FormulaRule(formula=[f'${ins}{FIRST}="Y"'], fill=GREEN))
    ws.conditional_formatting.add(f"{chk}{FIRST}:{chk}{LAST}", FormulaRule(formula=[f'LEFT(${chk}{FIRST},5)="CHECK"'], fill=RED))
    ws.conditional_formatting.add(f"{chk}{FIRST}:{chk}{LAST}", FormulaRule(formula=[f'LEFT(${chk}{FIRST},4)="NOTE"'], fill=AMBER))
    ws.row_dimensions[4].height = 58
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(CALC_COLUMNS))}{LAST}"
    ws.sheet_properties.tabColor = "7F7F7F"
    protect(ws, filter_=True)


# ----------------------------------------------------------------- Summary ----

KPI_FIRST = FIRST
TOP_SECTION = 32
CHAPTER_SECTION = TOP_SECTION + 13          # 45
ORIGIN_SECTION = CHAPTER_SECTION + CHAPTER_ROWS + 5   # 60


def build_summary(wb):
    ws = wb.create_sheet("Summary")
    title(ws, "Summary: what the action does to cost and margin (formulas only, except the yellow chapter list)",
          "Send nothing from here until B30 reads YES and the Review tab shows a checker sign-off for this version.")
    rng = lambda key: f"Calc!${CALC_COL[key]}${FIRST}:${CALC_COL[key]}${LAST}"  # noqa: E731
    ins = rng("inscope")
    headers(ws, 4, 1, ["Measure", "Value", "How it is calculated"])
    origin_total_row = ORIGIN_SECTION + 2 + ORIGIN_ROWS + 1
    chapter_other_row = CHAPTER_SECTION + 2 + CHAPTER_ROWS
    kpis = [
        ("Action ID", '=act_ID&""', "TIM Calc", "Action B5"),
        ("Action name", '=act_Name&""', "TIM Calc", "Action B6"),
        ("Importing country", '=act_Country&""', "TIM Calc", "Action B7"),
        ("Effective date", '=IF(ISNUMBER(act_EffDate),act_EffDate,"")', "TIM Calc Date", "Action B11 (shown, not applied)"),
        ("Default additional rate", '=IF(ISNUMBER(act_Rate),act_Rate,"none: rate lines only")', "TIM Calc Pct", "Action B15"),
        ("Rate lines on Action", "=COUNT(rate_Rate)", "TIM Calc Whole", "Rate lines with a rate (Action columns Y to AI)"),
        ("Brief ID and version", '=act_Brief&""', "TIM Calc", "Action B9"),
        ("Quantity period", '=qty_Period&""', "TIM Calc", "Items K1"),
        ("Items in model", f'=SUMPRODUCT(--({rng("item")}<>""))', "TIM Calc Whole", "Rows on Items with an item number"),
        ("Items in scope", f'=COUNTIF({ins},"Y")', "TIM Calc Whole", "Calc column L = Y"),
        ("Research determinations listed", '=SUMPRODUCT(--(det_Key<>""))', "TIM Calc Whole",
         "Rows on the Action determinations list; each must show OK (B27)"),
        ("Research determinations applied",
         f'=COUNTIF({rng("det")},"IN")&" IN, "&COUNTIF({rng("det")},"OUT")&" OUT; in-scope items with a displaced duty: "'
         f'&COUNTIFS({ins},"Y",{rng("displ")},">0")', "TIM Calc",
         "Calc columns K and U"),
        ("In-scope items with no quantity", f'=COUNTIFS({ins},"Y",{rng("qty")},0)', "TIM Calc Whole",
         "Listed, but they add 0 to the totals. Get a quantity or state why none"),
        ("In-scope items without a sell price", f'=COUNTIFS({ins},"Y",{rng("priced")},"N")', "TIM Calc Whole",
         "Counted in the impact (B20, B21), left out of revenue and margin (B22 to B25)"),
        ("Rows with a CHECK flag (Calc)", f'=SUMPRODUCT(--(LEFT({rng("check")},5)="CHECK"))', "TIM Calc Whole",
         "Must be 0 before review"),
        ('="Impact for the period ("&qty_Period&"), in-scope items"', f'=SUMIFS({rng("impact")},{ins},"Y")', "TIM Calc Money",
         "Sum of landed cost change x quantity, all in-scope items"),
        ("Of which: items without a sell price", f'=SUMIFS({rng("impact")},{ins},"Y",{rng("priced")},"N")', "TIM Calc Money",
         "Part of B20 that no margin figure covers"),
        ("Revenue for the period, in-scope items with a sell price", f'=SUMIFS({rng("rev")},{ins},"Y")', "TIM Calc Money",
         "Sell price x quantity"),
        ("Margin % before, in-scope items with a sell price", f'=IF(B22=0,"",SUMIFS({rng("mg_b")},{ins},"Y")/B22)',
         "TIM Calc Pct", "Revenue-weighted; unpriced items excluded from both sides"),
        ("Margin % after, in-scope items with a sell price", f'=IF(B22=0,"",SUMIFS({rng("mg_a")},{ins},"Y")/B22)',
         "TIM Calc Pct", "Revenue-weighted; assumes sell prices do not change"),
        ("Margin change (points)", '=IF(B22=0,"",B24-B23)', "TIM Calc Pct", "After minus before"),
        ("Action tab complete?",
         '=IF(AND(act_ID&""<>"",act_Country&""<>"",act_Brief&""<>"",ISNUMBER(act_EffDate),'
         'OR(ISNUMBER(act_Rate),COUNT(rate_Rate)>0),act_RateType&""<>"",act_USMCAExempt&""<>"",'
         'SUMPRODUCT(--(scope_HTS<>""))>0,SUMPRODUCT(--(scope_COO<>""))>0,'
         'OR(UPPER(TRIM(act_USMCAExempt&""))<>"PARTIAL",COUNTIF(rate_USMCA,"Y")>0)),"YES","NO")', "TIM Calc",
         "Required fields; a default rate or a rate line; one HTS prefix and one origin; PARTIAL needs a USMCA = Y rate line"),
        ("Determinations and rate lines check out?",
         '=IF(SUMPRODUCT(--(LEFT(det_Check,5)="CHECK"))+SUMPRODUCT(--(LEFT(rate_Check,5)="CHECK"))=0,"YES","NO")', "TIM Calc",
         "Every determination matched one Items row of this import country; no CHECK on Action columns W or AI"),
        ("Origin table reconciles?", f'=IF(ABS(C{origin_total_row}-B20)<0.005,"YES","NO")', "TIM Calc",
         f"Origin table total (C{origin_total_row}) equals B20; fails if an origin is listed twice"),
        ("Chapter table reconciles?", f'=IF(ABS(C{chapter_other_row})<0.005,"YES","NO: add the missing chapter")', "TIM Calc",
         f"Nothing left in 'Other chapters' (C{chapter_other_row})"),
        ("Ready for checker review?", '=IF(AND(B26="YES",B27="YES",B28="YES",B29="YES",B19=0),"YES","NO")', "TIM Calc",
         "B26 to B29 all YES and B19 = 0"),
    ]
    for i, (label, formula, style, how) in enumerate(kpis):
        r = KPI_FIRST + i
        put(ws, f"A{r}", label, "TIM Label")
        put(ws, f"B{r}", formula, style)
        put(ws, f"C{r}", how, "TIM Note")
    assert KPI_FIRST + len(kpis) - 1 == 30 and TOP_SECTION > 31
    for ref in ("B26", "B27", "B28", "B29", "B30"):
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{ref}="YES"'], fill=GREEN))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},2)="NO"'], fill=RED))
    ws.conditional_formatting.add("B19", FormulaRule(formula=["B19>0"], fill=RED))

    # Top 10 by absolute impact (ties broken by row, so every item appears once).
    put(ws, f"A{TOP_SECTION}", "Top 10 in-scope items by size of impact", "TIM Section")
    headers(ws, TOP_SECTION + 1, 1, ["Rank", "Item", "Description", "Origin", "Rate line", "Impact for the period",
                                     "Change per unit", "Margin % after", "Rank key (helper)"])
    key_rng = rng("rank")
    for k in range(1, 11):
        r = TOP_SECTION + 1 + k
        put(ws, f"A{r}", k, "TIM Calc Whole")
        put(ws, f"I{r}", f'=IF(COUNT({key_rng})<A{r},"",LARGE({key_rng},A{r}))', "TIM Calc Key")
        for col, src, style in (("B", "item", "TIM Calc"), ("C", "desc", "TIM Calc"), ("D", "coo", "TIM Calc"),
                                ("E", "line", "TIM Calc"), ("F", "impact", "TIM Calc Money"),
                                ("G", "chg", "TIM Calc Unit"), ("H", "m_a", "TIM Calc Pct")):
            put(ws, f"{col}{r}", f'=IF($I{r}="","",INDEX({rng(src)},MATCH($I{r},{key_rng},0)))', style)

    # By HTS chapter: yellow list of chapters (text), plus an 'Other' row that must be zero.
    put(ws, f"A{CHAPTER_SECTION}", "By HTS chapter (in-scope items). Type each chapter in scope as 2 digits", "TIM Section")
    headers(ws, CHAPTER_SECTION + 1, 1, ["HTS chapter (2 digits)", "Items in scope", "Impact for the period", "Share of impact"])
    ch_first, ch_last = CHAPTER_SECTION + 2, CHAPTER_SECTION + 1 + CHAPTER_ROWS
    for r in range(ch_first, ch_last + 1):
        put(ws, f"A{r}", None, "TIM Input Text")
        put(ws, f"B{r}", f'=IF(A{r}="","",COUNTIFS({rng("chap")},A{r}&"",{ins},"Y"))', "TIM Calc Whole")
        put(ws, f"C{r}", f'=IF(A{r}="","",SUMIFS({rng("impact")},{rng("chap")},A{r}&"",{ins},"Y"))', "TIM Calc Money")
        put(ws, f"D{r}", f'=IF(OR(A{r}="",$B$20=0),"",C{r}/$B$20)', "TIM Calc Pct")
    ws[f"A{ch_first}"].value = "73"
    o = chapter_other_row
    assert o == ch_last + 1
    put(ws, f"A{o}", "Other chapters (must be 0)", "TIM Label")
    put(ws, f"B{o}", f"=B14-SUM(B{ch_first}:B{ch_last})", "TIM Calc Whole")
    put(ws, f"C{o}", f"=B20-SUM(C{ch_first}:C{ch_last})", "TIM Calc Money")
    put(ws, f"D{o}", f'=IF($B$20=0,"",C{o}/$B$20)', "TIM Calc Pct")
    ws.conditional_formatting.add(f"B{o}:C{o}", FormulaRule(formula=[f"ABS($C${o})>=0.005"], fill=RED))
    length_dv(ws, f"A{ch_first}:A{ch_last}", 2, "Type the chapter as 2 digits, for example 73 or 08.")

    # By country of origin: one row per row of the Action origin list, plus a directly computed 'Other' row.
    put(ws, f"A{ORIGIN_SECTION}", "By country of origin (in-scope items; rows follow the Action origin list)", "TIM Section")
    headers(ws, ORIGIN_SECTION + 1, 1, ["Origin (from Action list)", "Items in scope", "Impact for the period", "Share of impact"])
    or_first, or_last = ORIGIN_SECTION + 2, ORIGIN_SECTION + 1 + ORIGIN_ROWS
    for i in range(ORIGIN_ROWS):
        r = or_first + i
        put(ws, f"A{r}", f'=IF(Action!$M${FIRST + i}="","",UPPER(TRIM(Action!$M${FIRST + i}&"")))', "TIM Calc")
        put(ws, f"B{r}", f'=IF(A{r}="","",COUNTIFS({rng("coo")},A{r},{ins},"Y"))', "TIM Calc Whole")
        put(ws, f"C{r}", f'=IF(A{r}="","",SUMIFS({rng("impact")},{rng("coo")},A{r},{ins},"Y"))', "TIM Calc Money")
        put(ws, f"D{r}", f'=IF(OR(A{r}="",$B$20=0),"",C{r}/$B$20)', "TIM Calc Pct")
    other = or_last + 1
    assert other + 1 == origin_total_row
    put(ws, f"A{other}", "Other origins (in scope by an IN determination; origin not on the list)", "TIM Label")
    put(ws, f"B{other}", f'=COUNTIFS({ins},"Y",{rng("coo_ok")},"N")', "TIM Calc Whole")
    put(ws, f"C{other}", f'=SUMIFS({rng("impact")},{ins},"Y",{rng("coo_ok")},"N")', "TIM Calc Money")
    put(ws, f"D{other}", f'=IF($B$20=0,"",C{other}/$B$20)', "TIM Calc Pct")
    t = origin_total_row
    put(ws, f"A{t}", "Total (must equal B20)", "TIM Label")
    put(ws, f"B{t}", f"=SUM(B{or_first}:B{other})", "TIM Calc Whole")
    put(ws, f"C{t}", f"=SUM(C{or_first}:C{other})", "TIM Calc Money")
    put(ws, f"D{t}", f'=IF($B$20=0,"",C{t}/$B$20)', "TIM Calc Pct")
    ws.conditional_formatting.add(f"B{t}:C{t}", FormulaRule(formula=[f"ABS($C${t}-$B$20)>=0.005"], fill=RED))

    widths(ws, {"A": 46, "B": 22, "C": 46, "D": 14, "E": 13, "F": 14, "G": 13, "H": 12, "I": 13})
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = "548235"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    protect(ws)
    return {"top_first": TOP_SECTION + 2, "chapter_first": ch_first, "chapter_other": o,
            "origin_first": or_first, "origin_other": other, "origin_total": t}


# ------------------------------------------------------------------ Review ----

REVIEW_HEADERS = ["Version", "Date", "Maker (name, track, what was prepared)", "Checker", "What the checker compared",
                  "Findings", "Findings closed? (Y/N)", "Decision", "Released to", "Release date"]
SAMPLE_REVIEWS = [
    ("v1.0 (SAMPLE)", dt.date(2027, 2, 18), "Hire A, Research: Action tab", "Senior",
     "Action tab against brief v1.0: every HTS prefix, exclusion, origin, rate, base, date and determination",
     "SAMPLE: stacking note reworded to quote the brief; closed", "Y", "Approved for release", "Standard Cost (Hire B)",
     dt.date(2027, 2, 18)),
    ("v1.1 (SAMPLE)", dt.date(2027, 2, 19), "Hire B, Standard Cost: Items, Summary", "Manager",
     "Item count and totals against the extract; units of measure; 3 items recalculated by hand; Summary B30 = YES",
     "SAMPLE: none", "Y", "Approved for release", "[COMPANY: Pricing lead], Finance", dt.date(2027, 2, 19)),
]
CHECKLIST = [
    "Version on this tab = Action B24 = file name",
    "Brief version on Action matches the checked brief",
    "Every HTS prefix, rate group, exclusion, origin and determination traced to the brief",
    "Default rate, rate lines, rate type, duty base, displaced duties and USMCA rule match brief Part 7, 6.5 and 6.6",
    "Item count on Items = extract row count; one row per item and import country",
    "Items G is the Oracle primary UOM; value, quantity, cost and price are all per that unit",
    "Dutiable shares (Items M) traced to their source",
    "No CHECK flags (Summary B19 = 0, B27 = YES); every NOTE explained",
    "3 items recalculated by hand to $0.01",
    "Top 10 items make business sense",
    "Origin and chapter tables reconcile",
    "All tabs still protected; no typed numbers over grey cells",
    "Out-of-scope reasons spot-checked (5 rows)",
    "Limits stated in the covering note",
]


def build_review(wb):
    ws = wb.create_sheet("Review")
    title(ws, "Review: maker-checker sign-off log",
          "One line per review. The maker fills columns A-C; the checker fills D-H. Nothing is released without a line here.",
          "SAMPLE lines below are FICTIONAL. Clear them before real use.")
    headers(ws, 4, 1, REVIEW_HEADERS)
    styles = ["TIM Input Text", "TIM Input Date", "TIM Input", "TIM Input", "TIM Input", "TIM Input",
              "TIM Input Text", "TIM Input", "TIM Input", "TIM Input Date"]
    for r in range(FIRST, 45):
        for c, style in enumerate(styles, 1):
            put(ws, f"{get_column_letter(c)}{r}", None, style)
    for i, row in enumerate(SAMPLE_REVIEWS):
        for c, v in enumerate(row, 1):
            ws.cell(row=FIRST + i, column=c).value = v
    list_dv(ws, ["Y", "N"], "G5:G44")
    list_dv(ws, ["Approved for release", "Approved with changes", "Returned to maker"], "H5:H44")
    number_dv(ws, "date", "B5:B44", 36526, 73051, "Enter a date (yyyy-mm-dd).")
    number_dv(ws, "date", "J5:J44", 36526, 73051, "Enter a date (yyyy-mm-dd).")

    headers(ws, 4, 12, ["Checker checklist (each review)", "Done (Y/N)"])
    for i, item in enumerate(CHECKLIST):
        put(ws, f"L{FIRST + i}", item, "TIM Text")
        put(ws, f"M{FIRST + i}", None, "TIM Input Text")
    list_dv(ws, ["Y", "N"], f"M{FIRST}:M{FIRST + len(CHECKLIST) - 1}")
    widths(ws, {"A": 13, "B": 11, "C": 30, "D": 12, "E": 40, "F": 34, "G": 10, "H": 20, "I": 22, "J": 11,
                "K": 2, "L": 50, "M": 10})
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = "7030A0"
    protect(ws)


# ------------------------------------------------------------------- write ----

FIXED_ZIP_TIME = (2026, 10, 8, 0, 0, 0)
FIXED_ISO = TEMPLATE_DATE.strftime("%Y-%m-%dT%H:%M:%SZ")


def make_deterministic(src, dst):
    """Rewrite the xlsx zip with fixed timestamps so identical input gives identical bytes."""
    with zipfile.ZipFile(src) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    dst.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(dst, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for fname, data in entries:
            if fname == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                              lambda m: m.group(1) + FIXED_ISO + m.group(3), text)
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(fname, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, data)


def build(out):
    wb = Workbook()
    add_styles(wb)
    build_readme(wb)
    build_action(wb, CALC_COL)
    build_items(wb)
    build_calc(wb)
    build_summary(wb)
    build_review(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "BSA Manager"
    wb.properties.title = "Tariff-Impact Model (template)"
    wb.properties.created = TEMPLATE_DATE
    wb.properties.modified = TEMPLATE_DATE
    wb.properties.lastModifiedBy = "build/make_tariff_model.py"
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.xlsx"
        wb.save(raw)
        make_deterministic(raw, out)
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output .xlsx path")
    args = parser.parse_args()
    out = build(args.out.resolve())
    try:
        shown = out.relative_to(ROOT)
    except ValueError:
        shown = out
    print(f"tariff model: {shown}")


if __name__ == "__main__":
    main()
