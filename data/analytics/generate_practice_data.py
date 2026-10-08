#!/usr/bin/env python3
"""Generate the SYNTHETIC practice data for the BSA analytics practice pack (module M3).

Run from the repo root:

    python3 data/analytics/generate_practice_data.py

Writes:
    dist/practice-data/*.csv                     one CSV per table (UTF-8, comma-delimited, header row)
    dist/practice-data/load_practice_data.sql    T-SQL that creates schema `practice` and loads every table
    dist/practice-data/README.txt                what the data is and how to load it
    dist/Practice-Data.xlsx                      one sheet (and one Excel table) per table, plus a README sheet

Used by docs/analytics/practice-exercises.md. The answers are in docs/analytics/answer-key.md.
TRAINER-ONLY until the day-30 review, like the answer key: the hires never get this file before then.
The comments below give away the data-quality findings and the answers to Exercises 1 to 6.

Deterministic: one fixed seed, no use of today's date, no iteration over sets. Every output file is
byte-identical on every run (the xlsx save time and zip entry times are pinned; see normalize_xlsx).
Needs only the Python standard library, pandas and openpyxl.

EVERYTHING HERE IS INVENTED. A fictional importer that brings hardware, electrical and household goods
into the US, Canada and Mexico. Item numbers, suppliers, brokers, entry numbers, values, costs and
duty rates are made up. HTS codes are plausible in format only: the first 6 digits are real HS
subheadings, the last 4 digits are random. Synthetic: do not use for classification. Rates are not
the real rates for these codes. All money is in USD to keep the practice simple.

Simplifications (say them out loud when you teach with this data):
    - One hts10 per item is used for all three import countries. In reality each country has its own
      10-digit code (US HTSUS, Canada Customs Tariff, Mexico TIGIE fraccion + NICO); only the first
      6 digits (HS) are shared.
    - HTS and origin come from the item master, not from the filed entry line.
    - A USMCA claim is valid in this data only when the item's country_of_origin is US, CA or MX AND
      usmca_eligible_flag = 'Y' (certification on file). In reality USMCA originating status comes
      from the product-specific rule of origin, not from the country-of-origin field.
    - A valid USMCA claim means 0% duty. No fees (MPF, HMF, DTA), no taxes (GST, IVA).
    - Receipts are receipts into the US inventory organization only, and std_cost is that org's cost.

PLANTED DATA-QUALITY ISSUES AND TRAPS (exact counts; generate() asserts every one of them):
    entry_lines duplicates ........ 24 extra rows: 4 US entries (6 lines each: US26-00110, US26-00114,
                                    US26-00163, US26-00177) were loaded twice. The copies are appended at
                                    the end of entry_lines.csv. Rows: 1,966 raw -> 1,942 after de-duplication.
    orphan entry lines ............ 4 lines whose item_id (SYN-9001, SYN-9002 x2, SYN-9003) is not in items.
    items with NULL hts10 ......... 6 items (3 CN, 2 VN, 1 IN); 94 clean entry lines use them.
    contradictory USMCA flag ...... 2 CN-origin items carry usmca_eligible_flag = 'Y'. All 30 of their clean
                                    lines claim USMCA (systematic invalid claims caused by master data).
    USMCA-origin items without
    a certification ............... 4 items (2 MX, 1 CA, 1 US) with usmca_eligible_flag = 'N'.
    applied rate <> reference ..... 49 clean non-preference lines (program_claimed NULL), planted as:
                                      18 stale rate: US, CN origin, ch. 73/76, entered on/after 2026-06-01
                                         at the old rate (14 by BRK-US2, 4 by BRK-US1)            -> under
                                       5 early rate: same goods entered in May 2026 at the new rate -> over
                                      11 additional duty missed: US, CN origin, other chapters      -> under
                                         (4 of them have a 0% base rate, so the data shows 0% applied)
                                       9 rate of another tariff line (classification question)   -> 2 over, 7 under
                                       6 0% applied on a dutiable line                              -> under
                                    Totals: 7 over-applied, 42 under-applied.
    invalid USMCA claims .......... 45 clean lines: 38 on non-USMCA origins (30 from the 2 flagged CN
                                    items + 8 one-off), 7 on USMCA-origin items with no certification.
    missed USMCA savings .......... 12 clean lines: eligible item, dutiable, no claim made.
    effective-dated reference ..... 11 US / CN / ch. 73-76 combinations have two periods (rate change on
                                    2026-06-01). A join without the date condition fans out: 107 clean
                                    non-preference lines match both periods (1,296 lines -> 1,403 rows).
    supplier country <> origin .... 24 items bought from a supplier in another country (traders, a distributor).
    outlier receipts .............. 12 of 1,500 receipts: 9 at about 10x or 12x standard (case cost booked
                                    per each), 3 at 0.1x standard (decimal shift). Rule: actual/std > 3 or < 1/3.
    overlapping scenario prefixes . SCN-01 lists 7318 (15%) and 731815 (25%) for CN; 20 clean lines match
                                    both. Rule: the longest matching prefix wins.
"""

import csv
import io
import random
import re
import sys
import zipfile
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

GENERATOR_VERSION = "1.1 (2026-10-08)"  # 1.1: same data rows as 1.0; reproducible xlsx
SEED = 20261019
DATA_FROM = date(2026, 1, 2)   # first entry date in the data
DATA_TO = date(2026, 9, 30)    # data cut-off: 9 months of history (annualize x 12/9)
STD_COST_DATE = date(2026, 1, 1)
RATE_CHANGE = date(2026, 6, 1)  # fictional US rate change for CN-origin chapters 73 and 76
OPEN_END = date(9999, 12, 31)
USMCA = ("US", "CA", "MX")

ROOT = Path(__file__).resolve().parent.parent.parent
OUT_DIR = ROOT / "dist" / "practice-data"
XLSX = ROOT / "dist" / "Practice-Data.xlsx"

D2 = Decimal("0.01")
D4 = Decimal("0.0001")


def q2(x):
    return Decimal(x).quantize(D2, ROUND_HALF_UP)


def q4(x):
    return Decimal(x).quantize(D4, ROUND_HALF_UP)


def dec(x):
    """Float to Decimal through a short string, so results do not depend on float repr details."""
    return Decimal(f"{x:.6f}")


# ------------------------------------------------------------------ reference data ----

# hs6: (chapter, description template, sizes, unit FOB range USD, qty-per-line range)
FAMILIES = {
    "731815": ("73", "Hex bolt, zinc-plated steel, {s}, box of 100", ["M6 x 30 mm", "M8 x 40 mm", "M10 x 50 mm", "M12 x 60 mm"], (6, 18), (200, 2400)),
    "731816": ("73", "Hex nut, steel, {s}, box of 200", ["M6", "M8", "M10", "M12"], (4, 11), (200, 2400)),
    "731822": ("73", "Flat washer, steel, {s}, box of 500", ["6 mm", "8 mm", "10 mm", "12 mm"], (3, 9), (200, 2000)),
    "732690": ("73", "Steel mounting bracket, {s}", ["50 mm", "75 mm", "100 mm", "150 mm"], (1.2, 4.5), (500, 6000)),
    "730722": ("73", "Stainless threaded elbow, {s}", ["1/4 in", "1/2 in", "3/4 in", "1 in"], (1.5, 6), (300, 4000)),
    "732393": ("73", "Stainless steel mixing bowl set, {s}", ["3-piece", "5-piece"], (6, 14), (100, 1200)),
    "761510": ("76", "Aluminum sheet pan, {s}", ["9 x 13 in", "13 x 18 in"], (2.5, 7), (200, 3000)),
    "761699": ("76", "Aluminum storage hook, {s}", ["small", "large", "double"], (0.8, 2.5), (500, 6000)),
    "820411": ("82", "Combination wrench set, {s}", ["8-piece", "12-piece", "16-piece"], (9, 30), (50, 800)),
    "820320": ("82", "Slip-joint pliers, {s}", ["6 in", "8 in", "10 in"], (2, 6), (200, 3000)),
    "820540": ("82", "Screwdriver set, {s}", ["6-piece", "10-piece"], (3, 10), (100, 1500)),
    "830110": ("83", "Padlock, laminated steel, {s}", ["40 mm", "50 mm"], (2.5, 7), (200, 2500)),
    "830210": ("83", "Butt hinge, steel, {s}", ["3 in", "3.5 in", "4 in"], (0.9, 3), (500, 5000)),
    "848180": ("84", "Ball valve, brass, {s}", ["1/2 in", "3/4 in", "1 in"], (4, 15), (100, 1500)),
    "848130": ("84", "Check valve, brass, {s}", ["1/2 in", "3/4 in", "1 in"], (5, 16), (100, 1200)),
    "841370": ("84", "Centrifugal utility pump, {s}", ["1/4 HP", "1/2 HP"], (35, 90), (20, 300)),
    "841451": ("84", "Oscillating floor fan, {s}", ["16 in", "18 in", "20 in"], (12, 28), (50, 800)),
    "846721": ("84", "Cordless drill, {s}", ["12 V", "18 V", "20 V"], (22, 60), (40, 600)),
    "850440": ("85", "Battery charger, {s}", ["12 V", "18 V", "20 V"], (7, 20), (100, 1500)),
    "850760": ("85", "Lithium-ion battery pack, {s}", ["2 Ah", "4 Ah", "5 Ah"], (14, 40), (50, 1000)),
    "853650": ("85", "Wall switch, {s}", ["15 A single-pole", "15 A 3-way", "20 A single-pole"], (0.6, 2.2), (500, 8000)),
    "853669": ("85", "Duplex receptacle, {s}", ["15 A", "20 A", "15 A tamper-resistant"], (0.7, 2.6), (500, 8000)),
    "854442": ("85", "Extension cord, {s}", ["10 ft", "25 ft", "50 ft"], (3, 12), (200, 3000)),
    "392690": ("39", "Plastic cable clip, {s}", ["pack of 50", "pack of 100"], (0.8, 3), (500, 6000)),
    "392310": ("39", "Plastic storage bin, {s}", ["12 L", "27 L", "45 L"], (2, 7), (200, 3000)),
    "391732": ("39", "Vinyl tubing, {s}", ["1/4 in x 20 ft", "3/8 in x 20 ft", "1/2 in x 10 ft"], (1.5, 6), (200, 3000)),
    "401693": ("40", "Rubber O-ring kit, {s}", ["225-piece", "407-piece"], (3, 9), (100, 1500)),
    "940320": ("94", "Steel shelving unit, {s}", ["4-tier", "5-tier"], (18, 45), (30, 500)),
    "630790": ("63", "Moving blanket, {s}", ["72 x 80 in", "40 x 72 in"], (3, 9), (100, 1500)),
    "420292": ("42", "Tool bag, polyester, {s}", ["13 in", "16 in", "20 in"], (4, 12), (100, 1500)),
}

# Items per country of origin, and the product families each origin supplies.
COO_PLAN = [
    ("CN", 48, None),  # None = every family
    ("VN", 16, ["731815", "731816", "731822", "848180", "848130", "392690", "392310", "420292", "630790", "940320", "830110", "820320"]),
    ("MX", 22, ["853650", "853669", "854442", "850440", "841451", "841370", "940320", "732690", "392310", "761510"]),
    ("US", 18, ["850440", "854442", "940320", "848180", "391732", "401693", "846721", "732690"]),
    ("CA", 14, ["761510", "761699", "732690", "940320", "391732", "401693", "841370"]),
    ("IN", 9, ["730722", "732393", "848180", "848130", "731815", "820411"]),
    ("TW", 8, ["731815", "731816", "820411", "820540", "846721", "850760"]),
    ("DE", 6, ["846721", "848180", "820411", "820540"]),
    ("KR", 5, ["850760", "853669", "854442"]),
    ("IT", 4, ["848180", "848130", "841370"]),
]

# Synthetic base (MFN-like) rates by import country. NOT real rates.
BASE_RATES = {
    "US": ["0.0000", "0.0250", "0.0290", "0.0340", "0.0390", "0.0450", "0.0570", "0.0650"],
    "CA": ["0.0000", "0.0300", "0.0600", "0.0650", "0.0800"],
    "MX": ["0.0000", "0.0500", "0.1000", "0.1500"],
}
# Synthetic additional duty, US imports of CN-origin goods. NOT a real tariff action.
US_CN_ADD = Decimal("0.1000")
US_CN_ADD_7376_NEW = Decimal("0.2500")  # chapters 73 and 76 from RATE_CHANGE

FREIGHT_PCT = {"CN": (0.07, 0.12), "VN": (0.07, 0.12), "IN": (0.08, 0.12), "TW": (0.06, 0.10), "KR": (0.06, 0.10),
               "DE": (0.05, 0.09), "IT": (0.05, 0.09), "MX": (0.03, 0.06), "CA": (0.03, 0.06), "US": (0.02, 0.04)}

SUPPLIER_PLAN = [  # (country, type, how many)
    ("CN", "Manufacturer", 6), ("VN", "Manufacturer", 3), ("MX", "Manufacturer", 4), ("US", "Manufacturer", 3),
    ("CA", "Manufacturer", 3), ("IN", "Manufacturer", 2), ("TW", "Manufacturer", 2), ("DE", "Manufacturer", 2),
    ("KR", "Manufacturer", 1), ("IT", "Manufacturer", 1), ("HK", "Trading company", 2), ("TW", "Trading company", 1),
    ("US", "Distributor", 1),
]
# Share of items per origin bought through a trading company or distributor, and which ones.
INDIRECT = {"CN": (0.25, [("HK", "Trading company"), ("TW", "Trading company")]),
            "VN": (0.15, [("TW", "Trading company")]),
            "MX": (0.15, [("US", "Distributor")]),
            "DE": (0.30, [("US", "Distributor")])}

ENTRIES_PER_COUNTRY = {"US": 230, "CA": 95, "MX": 75}
PORTS = {
    "US": ["Los Angeles, CA", "Long Beach, CA", "Savannah, GA", "Laredo, TX", "Detroit, MI", "Seattle, WA"],
    "CA": ["Vancouver, BC", "Toronto, ON", "Montreal, QC", "Windsor, ON"],
    "MX": ["Manzanillo, COL", "Nuevo Laredo, TAMPS", "Lazaro Cardenas, MICH", "Veracruz, VER"],
}
BROKERS = [  # (broker_id, broker_name, import_country, share of that country's entries)
    ("BRK-US1", "Fictional Customs Broker US-1", "US", 0.6), ("BRK-US2", "Fictional Customs Broker US-2", "US", 0.4),
    ("BRK-CA1", "Fictional Customs Broker CA-1", "CA", 0.7), ("BRK-CA2", "Fictional Customs Broker CA-2", "CA", 0.3),
    ("BRK-MX1", "Fictional Agente Aduanal MX-1", "MX", 0.6), ("BRK-MX2", "Fictional Agente Aduanal MX-2", "MX", 0.4),
]

# Fictional tariff scenarios for Exercise 6. NOT real actions.
SCENARIOS = [
    # scenario_id, name, import_country, prefix, origin, additional_rate, effective_date, usmca_exempt
    ("SCN-01", "Fictional US additional duty on fasteners and valves", "US", "7318", "CN", "0.1500", date(2026, 12, 1), "N"),
    ("SCN-01", "Fictional US additional duty on fasteners and valves", "US", "731815", "CN", "0.2500", date(2026, 12, 1), "N"),
    ("SCN-01", "Fictional US additional duty on fasteners and valves", "US", "8481", "CN", "0.1500", date(2026, 12, 1), "N"),
    ("SCN-01", "Fictional US additional duty on fasteners and valves", "US", "7318", "VN", "0.1000", date(2026, 12, 1), "N"),
    ("SCN-01", "Fictional US additional duty on fasteners and valves", "US", "8481", "VN", "0.1000", date(2026, 12, 1), "N"),
    ("SCN-02", "Fictional Canadian surtax on selected US goods", "CA", "8504", "US", "0.2500", date(2026, 11, 15), "N"),
    ("SCN-02", "Fictional Canadian surtax on selected US goods", "CA", "8544", "US", "0.2500", date(2026, 11, 15), "N"),
    ("SCN-02", "Fictional Canadian surtax on selected US goods", "CA", "9403", "US", "0.2500", date(2026, 11, 15), "N"),
    ("SCN-03", "Fictional US duty on Mexican electrical goods, USMCA-qualifying exempt", "US", "8536", "MX", "0.2500", date(2027, 1, 1), "Y"),
    ("SCN-03", "Fictional US duty on Mexican electrical goods, USMCA-qualifying exempt", "US", "8544", "MX", "0.2500", date(2027, 1, 1), "Y"),
    ("SCN-03", "Fictional US duty on Mexican electrical goods, USMCA-qualifying exempt", "US", "8504", "MX", "0.2500", date(2027, 1, 1), "Y"),
]

# Expected counts. generate() recomputes each one from the output tables and stops if any differs.
EXPECTED = {
    "items": 150, "suppliers": 31, "brokers": 6, "entries": 400, "entry_lines_raw": 1966, "entry_lines_clean": 1942,
    "duplicate_rows": 24, "duplicate_entries": 4, "orphan_lines": 4, "orphan_item_ids": 3,
    "null_hts_items": 6, "null_hts_clean_lines": 94, "contradictory_flag_items": 2, "uncertified_usmca_origin_items": 4,
    "rate_mismatch_lines": 49, "over_applied_lines": 7, "under_applied_lines": 42,
    "invalid_usmca_non_usmca_origin": 38, "invalid_usmca_no_certification": 7, "missed_usmca_lines": 12,
    "supplier_country_ne_origin_items": 24, "receipts": 1500, "outlier_receipts": 12,
    "scn01_overlap_lines": 20, "reference_rows": 228, "tariff_scenario_rows": 11,
}


# ----------------------------------------------------------------------- helpers ----

def business_days(start, end):
    days, d = [], start
    while d <= end:
        if d.weekday() < 5:
            days.append(d)
        d += timedelta(days=1)
    return days


def chapter(hts):
    return hts[:2]


class World:
    """Holds the synthetic base rates and the (fictional) additional-duty rules."""

    def __init__(self):
        self.base = {}

    def additional(self, ic, coo, hts, d):
        if ic == "US" and coo == "CN":
            if chapter(hts) in ("73", "76") and d >= RATE_CHANGE:
                return US_CN_ADD_7376_NEW
            return US_CN_ADD
        return Decimal("0")

    def expected(self, ic, hts, coo, d):
        return self.base[(ic, hts)] + self.additional(ic, coo, hts, d)


# --------------------------------------------------------------------- generation ----

def generate():
    rng = random.Random(SEED)
    world = World()
    bdays = business_days(DATA_FROM, DATA_TO)

    # HTS variants: real HS6 + 4 random digits (synthetic), 3 per family.
    hts_variants = {}
    for hs6 in FAMILIES:
        variants = []
        while len(variants) < 3:
            v = hs6 + rng.choice(["00", "20", "40", "50", "60", "80"]) + rng.choice(["10", "20", "30", "40", "50", "60", "80", "90"])
            if v not in variants:
                variants.append(v)
        hts_variants[hs6] = variants
    for hs6 in FAMILIES:
        for v in hts_variants[hs6]:
            for ic in ("US", "CA", "MX"):
                world.base[(ic, v)] = Decimal(rng.choice(BASE_RATES[ic]))

    # Suppliers
    suppliers, sup_by = [], {}
    n = 101
    for country, stype, count in SUPPLIER_PLAN:
        for k in range(count):
            label = {"Manufacturer": "Synthetic Supplier", "Trading company": "Synthetic Trading Co", "Distributor": "Synthetic Distributor"}[stype]
            sid = f"SUP-{n}"
            suppliers.append({"supplier_id": sid, "supplier_name": f"{label} {country}-{k + 1}", "supplier_country": country, "supplier_type": stype})
            sup_by.setdefault((country, stype), []).append(sid)
            n += 1

    # Items
    plan = []
    for coo, count, fams in COO_PLAN:
        pool = sorted(FAMILIES) if fams is None else list(fams)
        rng.shuffle(pool)
        for k in range(count):
            plan.append((coo, pool[k % len(pool)]))
    rng.shuffle(plan)
    items = []
    for k, (coo, hs6) in enumerate(plan):
        ch, tmpl, sizes, (lo, hi), qty_rng = FAMILIES[hs6]
        fob = q4(dec(rng.uniform(lo, hi)))
        frt = dec(rng.uniform(*FREIGHT_PCT[coo]))
        share, indirect = INDIRECT.get(coo, (0.0, []))
        if indirect and rng.random() < share:
            sid = rng.choice(sup_by[rng.choice(indirect)])
        else:
            sid = rng.choice(sup_by[(coo, "Manufacturer")])
        if coo == "US":
            ics = ["CA", "MX"] if rng.random() < 0.7 else ["CA"]
        elif coo in ("CA", "MX"):
            other = "MX" if coo == "CA" else "CA"
            ics = ["US", other] if rng.random() < 0.6 else ["US"]
        else:
            ics = ["US"] + (["CA"] if rng.random() < 0.45 else []) + (["MX"] if rng.random() < 0.35 else [])
        items.append({
            "item_id": f"SYN-{1001 + k}", "description": tmpl.format(s=rng.choice(sizes)), "hts10": rng.choice(hts_variants[hs6]),
            "country_of_origin": coo, "supplier_id": sid, "usmca_eligible_flag": "Y" if coo in USMCA else "N",
            "_hs6": hs6, "_fob": fob, "_frt": frt, "_ics": ics, "_qty": qty_rng, "_hts": None,
        })
    for it in items:
        it["_hts"] = it["hts10"]  # the true code, kept even when the master value is blanked below

    # Plant: USMCA-origin items with no certification (2 MX electrical, 1 CA, 1 US).
    def pick(cond, k):
        cands = [it for it in items if cond(it)]
        return rng.sample(cands, k)
    elec = ("853650", "853669", "854442", "850440")
    uncert = (pick(lambda it: it["country_of_origin"] == "MX" and it["_hs6"] in elec, 2)
              + pick(lambda it: it["country_of_origin"] == "CA", 1) + pick(lambda it: it["country_of_origin"] == "US", 1))
    for it in uncert:
        it["usmca_eligible_flag"] = "N"
    # Plant: contradictory master data, CN origin flagged USMCA-eligible.
    contra = pick(lambda it: it["country_of_origin"] == "CN" and it["_hs6"] in ("848180", "853669", "392690", "820320", "830210"), 2)
    for it in contra:
        it["usmca_eligible_flag"] = "Y"
    # Plant: NULL hts10 on 6 non-USMCA items.
    taken = {it["item_id"] for it in contra}
    nullh = (pick(lambda it: it["country_of_origin"] == "CN" and it["item_id"] not in taken and it["_hs6"][:2] not in ("73", "76"), 3)
             + pick(lambda it: it["country_of_origin"] == "VN", 2) + pick(lambda it: it["country_of_origin"] == "IN", 1))
    for it in nullh:
        it["hts10"] = None

    # Supplier price changes during 2026 (not on CN ch. 73/76 items, so the tariff story stays clean).
    price_change = {}
    eligible_pc = [it for it in items if not (it["country_of_origin"] == "CN" and it["_hs6"][:2] in ("73", "76"))]
    chosen = rng.sample(eligible_pc, 16)
    for k, it in enumerate(chosen):
        month = rng.randint(3, 7)
        factor = dec(rng.uniform(1.06, 1.15)) if k < 10 else dec(rng.uniform(0.90, 0.95))
        price_change[it["item_id"]] = (date(2026, month, 1), factor)

    def fob_on(it, d):
        pc = price_change.get(it["item_id"])
        return it["_fob"] * (pc[1] if pc and d >= pc[0] else Decimal(1))

    def us_duty_paid(it, d):
        if it["country_of_origin"] == "US" or it["usmca_eligible_flag"] == "Y":
            return Decimal(0)
        return world.expected("US", it["_hts"], it["country_of_origin"], d)

    for it in items:
        std = q4(it["_fob"] * (1 + it["_frt"] + us_duty_paid(it, STD_COST_DATE)))
        it["std_cost"] = std
        it["list_price"] = q2(std * dec(rng.uniform(1.35, 1.85)))

    # Entries
    entries = []
    for ic, count in ENTRIES_PER_COUNTRY.items():
        brk = [(b[0], b[3]) for b in BROKERS if b[2] == ic]
        rows = []
        for _ in range(count):
            r = rng.random()
            bid = brk[0][0] if r < brk[0][1] else brk[1][0]
            rows.append({"import_country": ic, "entry_date": rng.choice(bdays), "port": rng.choice(PORTS[ic]), "broker_id": bid})
        rows.sort(key=lambda e: e["entry_date"])
        for k, e in enumerate(rows):
            e["entry_no"] = f"{ic}26-{k + 1:05d}"
        entries.extend(rows)
    entries.sort(key=lambda e: (e["entry_date"], e["entry_no"]))
    ent_by = {e["entry_no"]: e for e in entries}

    # Entry lines (correct first; problems planted afterwards)
    pools = {ic: [it for it in items if ic in it["_ics"]] for ic in ("US", "CA", "MX")}
    lines = []
    for e in entries:
        k = rng.choice([2, 3, 4, 4, 5, 5, 6, 7, 8])
        chosen_items = rng.sample(pools[e["import_country"]], k)
        for ln, it in enumerate(chosen_items, 1):
            qlo, qhi = it["_qty"]
            qty = max(10, int(round(rng.randint(qlo, qhi), -1)))
            unit = fob_on(it, e["entry_date"]) * dec(rng.uniform(0.98, 1.02))
            value = q2(qty * unit)
            exp = world.expected(e["import_country"], it["_hts"], it["country_of_origin"], e["entry_date"])
            if it["usmca_eligible_flag"] == "Y":
                prog, rate = "USMCA", Decimal(0)
            else:
                prog, rate = None, exp
            lines.append({"entry_no": e["entry_no"], "line_no": ln, "item_id": it["item_id"], "qty": qty, "entered_value": value,
                          "duty_rate_applied": rate, "program_claimed": prog,
                          "_ic": e["import_country"], "_date": e["entry_date"], "_broker": e["broker_id"], "_exp": exp,
                          "_item": it, "_plant": None})

    # Duplicates: 4 US entries with exactly 6 plain lines are loaded twice.
    def plain(ln):
        it = ln["_item"]
        return it["hts10"] is not None and it not in contra and it not in uncert
    lines_by_entry = {}
    for ln in lines:
        lines_by_entry.setdefault(ln["entry_no"], []).append(ln)
    dup_cands = [e["entry_no"] for e in entries if e["import_country"] == "US"
                 and len(lines_by_entry[e["entry_no"]]) == 6 and all(plain(x) for x in lines_by_entry[e["entry_no"]])]
    dup_entries = rng.sample(dup_cands, 4)
    dup_set = set(dup_entries)

    def free(ln):
        return ln["_plant"] is None and ln["entry_no"] not in dup_set

    def sample_lines(cond, k, label):
        cands = [ln for ln in lines if free(ln) and cond(ln)]
        if len(cands) < k:
            sys.exit(f"not enough candidate lines for {label}: {len(cands)} < {k}")
        picked = rng.sample(cands, k)
        for ln in picked:
            ln["_plant"] = label
        return picked

    def eligible(it):
        return it["usmca_eligible_flag"] == "Y" and it["country_of_origin"] in USMCA

    # Mark the systematic invalid claims (lines of the contradictory items) so nothing else is planted on them.
    for ln in lines:
        if ln["_item"] in contra and ln["entry_no"] not in dup_set:
            ln["_plant"] = "usmca_invalid_contra"

    for ln in sample_lines(lambda l: eligible(l["_item"]) and l["_exp"] > 0, 12, "usmca_missed"):
        ln["program_claimed"], ln["duty_rate_applied"] = None, ln["_exp"]
    for ln in sample_lines(lambda l: l["_item"] in uncert, 7, "usmca_invalid_nocert"):
        ln["program_claimed"], ln["duty_rate_applied"] = "USMCA", Decimal(0)
    nonpref_ok = lambda l: l["program_claimed"] is None and l["_item"]["hts10"] is not None and not eligible(l["_item"])
    for ln in sample_lines(lambda l: nonpref_ok(l) and l["_item"]["country_of_origin"] not in USMCA, 8, "usmca_invalid_adhoc"):
        ln["program_claimed"], ln["duty_rate_applied"] = "USMCA", Decimal(0)

    cn7376 = lambda l: nonpref_ok(l) and l["_ic"] == "US" and l["_item"]["country_of_origin"] == "CN" and chapter(l["_item"]["_hts"]) in ("73", "76")
    for ln in (sample_lines(lambda l: cn7376(l) and l["_date"] >= RATE_CHANGE and l["_broker"] == "BRK-US2", 14, "stale")
               + sample_lines(lambda l: cn7376(l) and l["_date"] >= RATE_CHANGE and l["_broker"] == "BRK-US1", 4, "stale")):
        ln["duty_rate_applied"] = world.base[("US", ln["_item"]["_hts"])] + US_CN_ADD
    for ln in sample_lines(lambda l: cn7376(l) and date(2026, 5, 1) <= l["_date"] < RATE_CHANGE, 5, "early"):
        ln["duty_rate_applied"] = world.base[("US", ln["_item"]["_hts"])] + US_CN_ADD_7376_NEW
    for ln in sample_lines(lambda l: nonpref_ok(l) and l["_ic"] == "US" and l["_item"]["country_of_origin"] == "CN"
                           and chapter(l["_item"]["_hts"]) not in ("73", "76"), 11, "missing_add"):
        ln["duty_rate_applied"] = world.base[("US", ln["_item"]["_hts"])]
    for ln in sample_lines(lambda l: nonpref_ok(l) and l["_exp"] > 0, 9, "wrong_line"):
        add = world.additional(ln["_ic"], ln["_item"]["country_of_origin"], ln["_item"]["_hts"], ln["_date"])
        options = [Decimal(r) + add for r in BASE_RATES[ln["_ic"]] if Decimal(r) + add != ln["_exp"] and Decimal(r) > 0]
        ln["duty_rate_applied"] = rng.choice(options)
    for ln in sample_lines(lambda l: nonpref_ok(l) and l["_exp"] > 0, 6, "free"):
        ln["duty_rate_applied"] = Decimal(0)

    # Orphan lines: item ids that are not in the item master.
    orphan_entries = rng.sample([e["entry_no"] for e in entries if e["import_country"] == "US" and e["entry_no"] not in dup_set], 4)
    for eno, iid in zip(orphan_entries, ["SYN-9001", "SYN-9002", "SYN-9002", "SYN-9003"]):
        e = ent_by[eno]
        qty = 10 * rng.randint(20, 150)
        value = q2(qty * dec(rng.uniform(2, 12)))
        lines.append({"entry_no": eno, "line_no": len(lines_by_entry[eno]) + 1, "item_id": iid, "qty": qty, "entered_value": value,
                      "duty_rate_applied": Decimal("0.0390"), "program_claimed": None, "_ic": "US", "_date": e["entry_date"],
                      "_broker": e["broker_id"], "_exp": None, "_item": None, "_plant": "orphan"})
    lines.sort(key=lambda l: (l["_date"], l["entry_no"], l["line_no"]))
    for ln in lines:
        ln["duty_amount"] = q2(ln["entered_value"] * ln["duty_rate_applied"])
    dup_rows = [dict(ln) for ln in lines if ln["entry_no"] in dup_set]
    all_lines = lines + dup_rows

    # Duty rates reference: one row per import country + hts10 + origin used (items with an hts10), effective-dated.
    combos = sorted({(ic, it["hts10"], it["country_of_origin"]) for it in items if it["hts10"] for ic in it["_ics"]})
    reference = []
    for ic, hts, coo in combos:
        base = world.base[(ic, hts)]
        if ic == "US" and coo == "CN" and chapter(hts) in ("73", "76"):
            periods = [(date(2026, 1, 1), RATE_CHANGE - timedelta(days=1), US_CN_ADD), (RATE_CHANGE, OPEN_END, US_CN_ADD_7376_NEW)]
        else:
            periods = [(date(2026, 1, 1), OPEN_END, world.additional(ic, coo, hts, date(2026, 1, 1)))]
        for f, t, add in periods:
            reference.append({"import_country": ic, "hts10": hts, "country_of_origin": coo, "effective_from": f, "effective_to": t,
                              "base_rate": base, "additional_rate": add, "expected_rate": base + add})

    # Receipts into the US org: 1,500 rows on POs of 1 to 3 lines.
    receipts, po_seq = [], 0
    while len(receipts) < 1500:
        po_seq += 1
        d = rng.choice(bdays)
        for it in rng.sample(items, rng.choice([1, 1, 2, 2, 3])):
            if len(receipts) >= 1500:
                break
            qlo, qhi = it["_qty"]
            qty = max(10, int(round(rng.randint(qlo // 2, qhi // 2), -1)))
            fob = fob_on(it, d) * dec(rng.uniform(0.97, 1.03))
            landed = fob * (1 + it["_frt"] * dec(rng.uniform(0.9, 1.1)) + us_duty_paid(it, d))
            receipts.append({"po_no": f"PO-26-{po_seq:05d}", "item_id": it["item_id"], "receipt_date": d, "qty": qty,
                             "actual_landed_unit_cost": q4(landed), "_outlier": None})
    receipts.sort(key=lambda r: (r["receipt_date"], r["po_no"], r["item_id"]))
    for k, r in enumerate(receipts):
        r["receipt_id"] = f"RCV-{k + 1:06d}"
    for k, r in enumerate(rng.sample(receipts, 12)):
        factor = Decimal(rng.choice([10, 12])) if k < 9 else Decimal("0.1")
        r["actual_landed_unit_cost"] = q4(r["actual_landed_unit_cost"] * factor)
        r["_outlier"] = "high" if k < 9 else "low"

    # ---------------------------------------------------------------- frames ----
    t_items = pd.DataFrame([{c: it[c] for c in ["item_id", "description", "hts10", "country_of_origin", "supplier_id", "std_cost", "list_price", "usmca_eligible_flag"]}
                            for it in sorted(items, key=lambda x: x["item_id"])])
    t_suppliers = pd.DataFrame(suppliers)
    t_brokers = pd.DataFrame([{"broker_id": b[0], "broker_name": b[1], "import_country": b[2]} for b in BROKERS])
    t_entries = pd.DataFrame([{c: e[c] for c in ["entry_no", "import_country", "entry_date", "port", "broker_id"]} for e in entries])
    t_lines = pd.DataFrame([{c: ln[c] for c in ["entry_no", "line_no", "item_id", "qty", "entered_value", "duty_rate_applied", "duty_amount", "program_claimed"]}
                            for ln in all_lines])
    t_lines["duty_rate_applied"] = t_lines["duty_rate_applied"].map(q4)
    t_ref = pd.DataFrame(reference)
    for c in ("base_rate", "additional_rate", "expected_rate"):
        t_ref[c] = t_ref[c].map(q4)
    t_receipts = pd.DataFrame([{c: r[c] for c in ["receipt_id", "po_no", "item_id", "receipt_date", "qty", "actual_landed_unit_cost"]} for r in receipts])
    t_scn = pd.DataFrame([dict(zip(["scenario_id", "scenario_name", "import_country", "hts_prefix", "country_of_origin", "additional_rate", "effective_date", "usmca_exempt"],
                                   [s[0], s[1], s[2], s[3], s[4], Decimal(s[5]), s[6], s[7]])) for s in SCENARIOS])
    tables = {"items": t_items, "suppliers": t_suppliers, "brokers": t_brokers, "entries": t_entries, "entry_lines": t_lines,
              "duty_rates_reference": t_ref, "receipts": t_receipts, "tariff_scenario": t_scn}

    check_counts(tables)
    return tables


# ------------------------------------------------------------------- self-checks ----

def check_counts(tables):
    """Recompute every planted count from the output tables, the way a learner would, and compare."""
    it, el, ref = tables["items"], tables["entry_lines"], tables["duty_rates_reference"]
    ent = tables["entries"]
    clean = el.drop_duplicates(subset=["entry_no", "line_no"]).copy()
    j = clean.merge(ent, on="entry_no", how="left").merge(it, on="item_id", how="left")
    got = {
        "items": len(it), "suppliers": len(tables["suppliers"]), "brokers": len(tables["brokers"]), "entries": len(ent),
        "entry_lines_raw": len(el), "entry_lines_clean": len(clean),
        "duplicate_rows": int(el.duplicated(subset=["entry_no", "line_no"]).sum()),
        "duplicate_entries": int(el[el.duplicated(subset=["entry_no", "line_no"], keep=False)]["entry_no"].nunique()),
        "orphan_lines": int((~clean["item_id"].isin(it["item_id"])).sum()),
        "orphan_item_ids": int(clean.loc[~clean["item_id"].isin(it["item_id"]), "item_id"].nunique()),
        "null_hts_items": int(it["hts10"].isna().sum()),
        "null_hts_clean_lines": int((j["hts10"].isna() & j["country_of_origin"].notna()).sum()),
        "contradictory_flag_items": int(((it["usmca_eligible_flag"] == "Y") & ~it["country_of_origin"].isin(USMCA)).sum()),
        "uncertified_usmca_origin_items": int(((it["usmca_eligible_flag"] == "N") & it["country_of_origin"].isin(USMCA)).sum()),
        "supplier_country_ne_origin_items": int((it.merge(tables["suppliers"], on="supplier_id")
                                                 .pipe(lambda d: d["supplier_country"] != d["country_of_origin"])).sum()),
        "receipts": len(tables["receipts"]), "tariff_scenario_rows": len(tables["tariff_scenario"]),
    }
    # Applied vs reference (effective-dated), non-preference lines only.
    np_ = j[j["program_claimed"].isna() & j["hts10"].notna()].copy()
    m = np_.merge(ref, on=["import_country", "hts10", "country_of_origin"], how="left")
    m = m[(m["entry_date"] >= m["effective_from"]) & (m["entry_date"] <= m["effective_to"])]
    if len(m) != len(np_):
        sys.exit(f"effective-dated join changed the row count: {len(np_)} -> {len(m)}")
    diff = m["duty_rate_applied"] - m["expected_rate"]
    got["rate_mismatch_lines"] = int((diff != 0).sum())
    got["over_applied_lines"] = int((diff > 0).sum())
    got["under_applied_lines"] = int((diff < 0).sum())
    u = j[j["program_claimed"] == "USMCA"]
    got["invalid_usmca_non_usmca_origin"] = int((~u["country_of_origin"].isin(USMCA)).sum())
    got["invalid_usmca_no_certification"] = int((u["country_of_origin"].isin(USMCA) & (u["usmca_eligible_flag"] == "N")).sum())
    elig = j[j["program_claimed"].isna() & j["country_of_origin"].isin(USMCA) & (j["usmca_eligible_flag"] == "Y")]
    elig = elig.merge(ref, on=["import_country", "hts10", "country_of_origin"])
    elig = elig[(elig["entry_date"] >= elig["effective_from"]) & (elig["entry_date"] <= elig["effective_to"])]
    got["missed_usmca_lines"] = int((elig["expected_rate"] > 0).sum())
    r = tables["receipts"].merge(it[["item_id", "std_cost"]], on="item_id")
    ratio = r["actual_landed_unit_cost"].astype(float) / r["std_cost"].astype(float)
    got["outlier_receipts"] = int(((ratio > 3) | (ratio < 1 / 3)).sum())
    normal = ratio[(ratio <= 3) & (ratio >= 1 / 3)]
    if normal.max() > 1.6 or normal.min() < 0.7:
        sys.exit(f"normal receipts drifted too far from standard: {normal.min():.3f} to {normal.max():.3f}")
    s = j[(j["import_country"] == "US") & (j["country_of_origin"] == "CN") & j["hts10"].notna()]
    got["scn01_overlap_lines"] = int(s["hts10"].str.startswith("731815").sum())
    got["reference_rows"] = len(ref)

    bad = {k: (EXPECTED[k], got.get(k)) for k in EXPECTED if EXPECTED[k] != got.get(k)}
    if bad:
        for k, (want, have) in bad.items():
            print(f"  count mismatch {k}: expected {want}, got {have}")
        sys.exit("planted-issue counts changed: update EXPECTED and the docstring only on purpose")
    for name, df in tables.items():
        for col in df.columns:
            if not all(not isinstance(v, str) or v.isascii() for v in df[col]):
                sys.exit(f"non-ASCII text in {name}.{col}: keep the CSVs plain ASCII for every loader")


# ------------------------------------------------------------------------ output ----

def fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, Decimal):
        return format(v, "f")
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v)


def write_csvs(tables):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, df in tables.items():
        with open(OUT_DIR / f"{name}.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(list(df.columns))
            for row in df.itertuples(index=False):
                w.writerow([fmt(v) for v in row])


SQL_TYPES = {
    "items": [("item_id", "VARCHAR(10) NOT NULL"), ("description", "VARCHAR(100) NOT NULL"), ("hts10", "VARCHAR(10) NULL"),
              ("country_of_origin", "CHAR(2) NOT NULL"), ("supplier_id", "VARCHAR(10) NOT NULL"), ("std_cost", "DECIMAL(18,4) NOT NULL"),
              ("list_price", "DECIMAL(18,2) NOT NULL"), ("usmca_eligible_flag", "CHAR(1) NOT NULL")],
    "suppliers": [("supplier_id", "VARCHAR(10) NOT NULL"), ("supplier_name", "VARCHAR(60) NOT NULL"), ("supplier_country", "CHAR(2) NOT NULL"),
                  ("supplier_type", "VARCHAR(20) NOT NULL")],
    "brokers": [("broker_id", "VARCHAR(10) NOT NULL"), ("broker_name", "VARCHAR(60) NOT NULL"), ("import_country", "CHAR(2) NOT NULL")],
    "entries": [("entry_no", "VARCHAR(12) NOT NULL"), ("import_country", "CHAR(2) NOT NULL"), ("entry_date", "DATE NOT NULL"),
                ("port", "VARCHAR(40) NOT NULL"), ("broker_id", "VARCHAR(10) NOT NULL")],
    "entry_lines": [("entry_no", "VARCHAR(12) NOT NULL"), ("line_no", "INT NOT NULL"), ("item_id", "VARCHAR(10) NOT NULL"), ("qty", "INT NOT NULL"),
                    ("entered_value", "DECIMAL(18,2) NOT NULL"), ("duty_rate_applied", "DECIMAL(9,4) NOT NULL"),
                    ("duty_amount", "DECIMAL(18,2) NOT NULL"), ("program_claimed", "VARCHAR(10) NULL")],
    "duty_rates_reference": [("import_country", "CHAR(2) NOT NULL"), ("hts10", "VARCHAR(10) NOT NULL"), ("country_of_origin", "CHAR(2) NOT NULL"),
                             ("effective_from", "DATE NOT NULL"), ("effective_to", "DATE NOT NULL"), ("base_rate", "DECIMAL(9,4) NOT NULL"),
                             ("additional_rate", "DECIMAL(9,4) NOT NULL"), ("expected_rate", "DECIMAL(9,4) NOT NULL")],
    "receipts": [("receipt_id", "VARCHAR(12) NOT NULL"), ("po_no", "VARCHAR(12) NOT NULL"), ("item_id", "VARCHAR(10) NOT NULL"),
                 ("receipt_date", "DATE NOT NULL"), ("qty", "INT NOT NULL"), ("actual_landed_unit_cost", "DECIMAL(18,4) NOT NULL")],
    "tariff_scenario": [("scenario_id", "VARCHAR(10) NOT NULL"), ("scenario_name", "VARCHAR(100) NOT NULL"), ("import_country", "CHAR(2) NOT NULL"),
                        ("hts_prefix", "VARCHAR(10) NOT NULL"), ("country_of_origin", "CHAR(2) NOT NULL"), ("additional_rate", "DECIMAL(9,4) NOT NULL"),
                        ("effective_date", "DATE NOT NULL"), ("usmca_exempt", "CHAR(1) NOT NULL")],
}
# Keys are declared on every table except entry_lines: the raw entry lines contain duplicates, so a
# primary key would make the load fail. Exercise 1 is where you find out why.
SQL_KEYS = {"items": "item_id", "suppliers": "supplier_id", "brokers": "broker_id", "entries": "entry_no",
            "duty_rates_reference": "import_country, hts10, country_of_origin, effective_from", "receipts": "receipt_id",
            "tariff_scenario": "scenario_id, hts_prefix, country_of_origin"}
NUMERIC_SQL = ("INT", "DECIMAL")


def sql_literal(v, sqltype):
    s = fmt(v)
    if s == "":
        return "NULL"
    if sqltype.startswith(NUMERIC_SQL):
        return s
    return "'" + s.replace("'", "''") + "'"


def write_sql(tables):
    out = ["-- SYNTHETIC practice data for docs/analytics/practice-exercises.md. Do not use for classification.",
           f"-- Generated by the team's practice-data script (trainer-only), version {GENERATOR_VERSION}, seed {SEED}.",
           "-- Run in a scratch or sandbox database where you may create tables (never in a production database).",
           "-- It drops and recreates the schema `practice`, so you can rerun it to reset the data.",
           "SET NOCOUNT ON;", "GO"]
    for name in reversed(list(tables)):
        out.append(f"IF OBJECT_ID('practice.{name}', 'U') IS NOT NULL DROP TABLE practice.{name};")
    out += ["GO", "IF SCHEMA_ID('practice') IS NULL EXEC('CREATE SCHEMA practice');", "GO"]
    for name, df in tables.items():
        cols = SQL_TYPES[name]
        if [c for c, _ in cols] != list(df.columns):
            sys.exit(f"SQL column list out of step for {name}")
        body = ",\n    ".join(f"{c} {t}" for c, t in cols)
        key = f",\n    CONSTRAINT pk_{name} PRIMARY KEY ({SQL_KEYS[name]})" if name in SQL_KEYS else ""
        out.append(f"CREATE TABLE practice.{name} (\n    {body}{key}\n);")
        rows = [row for row in df.itertuples(index=False)]
        for start in range(0, len(rows), 500):
            chunk = rows[start:start + 500]
            values = ",\n".join("(" + ", ".join(sql_literal(v, t) for v, (_, t) in zip(row, cols)) + ")" for row in chunk)
            out.append(f"INSERT INTO practice.{name} ({', '.join(c for c, _ in cols)}) VALUES\n{values};")
        out.append("GO")
    out.append("-- Load check: compare with the row counts in practice-exercises.md (Setup).")
    out.append(" UNION ALL\n".join(f"SELECT '{n}' AS table_name, COUNT(*) AS row_count FROM practice.{n}" for n in tables) + ";")
    (OUT_DIR / "load_practice_data.sql").write_text("\n".join(out) + "\n", encoding="utf-8")


README_LINES = [
    "PRACTICE DATA: SYNTHETIC. Invented for training. Not real items, suppliers, brokers, entries, values, costs or rates.",
    "HTS codes look real (first 6 digits are real HS subheadings, last 4 are random). Synthetic: do not use for classification.",
    "Duty rates and tariff scenarios are invented. They are not the real rates or actions for these codes.",
    f"Generator: the team's practice-data script (trainer-only), version {GENERATOR_VERSION}, seed {SEED}. Same output on every run.",
    f"Entry and receipt dates: {DATA_FROM:%Y-%m-%d} to {DATA_TO:%Y-%m-%d} (9 months). All money is USD.",
    "Exercises: docs/analytics/practice-exercises.md (Analytics Practice Pack in OneNote). The answer key is trainer-only: the Senior reviews your results against it.",
    "",
    "TABLES (one row = ...; key)",
    "items: one item; item_id. hts10 and country_of_origin are the classification master. usmca_eligible_flag Y = USMCA certification on file. std_cost = US org standard cost per unit (includes duty and freight). list_price = sell price per unit.",
    "suppliers: one supplier; supplier_id. supplier_country is where the supplier is, not where the goods are made.",
    "brokers: one customs broker; broker_id.",
    "entries: one customs entry (US entry, Canadian CAD, Mexican pedimento); entry_no. import_country = country the goods entered.",
    "entry_lines: one line on an entry, as the broker filed it; entry_no + line_no (meant to be unique). program_claimed = USMCA, or blank when no preference was claimed. duty_amount = entered_value x duty_rate_applied, rounded to the cent.",
    "duty_rates_reference: the rate we expect for an import country + hts10 + origin, effective-dated (effective_from to effective_to, inclusive); key = those 3 columns + effective_from. expected_rate = base_rate + additional_rate. It is the non-preferential rate: a valid USMCA claim in this data means 0%.",
    "receipts: one PO receipt line into the US warehouse; receipt_id. actual_landed_unit_cost = actual cost per unit including freight and duty.",
    "tariff_scenario: one row per scenario + HTS prefix + origin. Fictional proposed actions for Exercise 6. If several prefixes in one scenario match an item, the longest prefix wins. usmca_exempt Y = items with usmca_eligible_flag Y are exempt.",
    "",
    "RATES are decimals: 0.0340 means 3.4%.",
    "LOADING: Excel/Power Query: use Practice-Data.xlsx (one Excel table per sheet). SQL Server: run load_practice_data.sql in a scratch database, or import the CSVs (see the exercises, Setup).",
    "Keep hts10, item_id and other codes as TEXT. If Excel turns hts10 into a number, re-import and set the column type to Text.",
    "The data contains deliberate data-quality problems, like real data. Finding them is part of the exercises.",
]


HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
MONEY2 = {"entered_value", "duty_amount", "list_price"}
MONEY4 = {"std_cost", "actual_landed_unit_cost"}
RATES = {"duty_rate_applied", "base_rate", "additional_rate", "expected_rate"}
DATES = {"entry_date", "effective_from", "effective_to", "receipt_date", "effective_date"}


def to_cell(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, Decimal):
        return float(v)
    return v


def write_xlsx(tables):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "Practice Data (SYNTHETIC) for the BSA analytics practice pack"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F3864")
    ws["A2"] = "SYNTHETIC: do not use for classification, rates or any business decision."
    ws["A2"].font = Font(name="Calibri", bold=True, color="C00000")
    for k, line in enumerate(README_LINES, 4):
        ws.cell(row=k, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    row = len(README_LINES) + 6
    ws.cell(row=row, column=1, value="Sheet, rows, columns").font = Font(bold=True)
    for name, df in tables.items():
        row += 1
        ws.cell(row=row, column=1, value=f"{name}: {len(df):,} rows; columns: {', '.join(df.columns)}").alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 150

    for name, df in tables.items():
        ws = wb.create_sheet(name)
        cols = list(df.columns)
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for rec in df.itertuples(index=False):
            ws.append([to_cell(v) for v in rec])
        for c, col in enumerate(cols, 1):
            letter = get_column_letter(c)
            is_text = all(isinstance(v, str) for v in df[col] if v is not None and not (isinstance(v, float) and pd.isna(v)))
            fmt_code = ("#,##0.00" if col in MONEY2 else "#,##0.0000" if col in MONEY4 else "0.0000" if col in RATES
                        else "yyyy-mm-dd" if col in DATES else "@" if is_text else None)
            if fmt_code:
                for r in range(2, len(df) + 2):
                    ws[f"{letter}{r}"].number_format = fmt_code
            width = max(len(col), *(len(fmt(v)) for v in df[col].head(200))) + 2
            ws.column_dimensions[letter].width = min(max(width, 10), 60)
        ws.freeze_panes = "A2"
        tab = Table(displayName=f"tbl_{name}", ref=f"A1:{get_column_letter(len(cols))}{len(df) + 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    fixed = datetime(2026, 10, 8)
    wb.properties.creator = "generate_practice_data.py"
    wb.properties.created = fixed
    wb.properties.modified = fixed
    buf = io.BytesIO()
    wb.save(buf)
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    XLSX.write_bytes(normalize_xlsx(buf.getvalue()))


XLSX_STAMP = "2026-10-08T00:00:00Z"
XLSX_ZIP_TIME = (2026, 10, 8, 0, 0, 0)


def normalize_xlsx(data):
    """openpyxl writes the save time into docProps/core.xml (dcterms:modified) and into every zip
    entry. Pin both so that reruns give a byte-identical file. The workbook content is unchanged."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:)",
                              rb"\g<1>" + XLSX_STAMP.encode() + rb"\g<2>", body)
            entry = zipfile.ZipInfo(info.filename, date_time=XLSX_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            dst.writestr(entry, body)
    return out.getvalue()


def main():
    tables = generate()
    write_csvs(tables)
    write_sql(tables)
    (OUT_DIR / "README.txt").write_text("\n".join(README_LINES) + "\n", encoding="utf-8")
    write_xlsx(tables)
    total = sum(len(df) for df in tables.values())
    print(f"practice data: {len(tables)} tables, {total:,} rows -> {OUT_DIR.relative_to(ROOT)}/ and {XLSX.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
